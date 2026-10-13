"""Exports the exception report (non-MATCHED reconciliation results) as XLSX."""
import io
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models import ReconciliationResult, MatchStatus
from app.money import to_float

HEADERS = [
    "Match Status", "Confidence", "Vendor Name", "Vendor GSTIN",
    "Invoice # (Books)", "Invoice # (GSTR-2B)", "Invoice Date (Books)",
    "Invoice Date (GSTR-2B)", "Taxable Value (Books)", "Taxable Value (GSTR-2B)",
    "Difference", "Notes", "Exception Status", "Assigned To", "Resolution Note",
    "Approved By", "ITC Lapse Deadline",
]

STATUS_COLORS = {
    MatchStatus.SUPPLIER_NOT_FILED.value: "FFF4CC",
    MatchStatus.VALUE_MISMATCH.value: "FCE4E4",
    MatchStatus.CREDIT_NOTE_DISCREPANCY.value: "FCE4E4",
    MatchStatus.MISSING_IN_BOOKS.value: "E4ECFC",
    MatchStatus.FORMATTING_MISMATCH.value: "E8F5E9",
}


def build_exception_report_xlsx(results: List[ReconciliationResult], client_name: str, period: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Exceptions"

    ws.append([f"GSTR-2B Reconciliation — Exception Report"])
    ws.append([f"Client: {client_name}   Period: {period}"])
    ws.append([])
    header_row_idx = 4
    ws.append(HEADERS)
    for cell in ws[header_row_idx]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")

    row_idx = header_row_idx + 1
    for r in results:
        p = r.purchase_entry
        g = r.gstr2b_entry
        exc = getattr(r, "exception", None)
        # Coerce Decimal -> float so openpyxl writes a plain number cell
        # and so the difference below never mixes Decimal with float.
        book_val = to_float(p.taxable_value) if p else None
        gstr_val = to_float(g.taxable_value) if g else None
        diff = round(book_val - gstr_val, 2) if (book_val is not None and gstr_val is not None) else None

        ws.append([
            r.match_status.value if hasattr(r.match_status, "value") else r.match_status,
            r.confidence_score,
            (p.vendor_name if p else (g.vendor_name if g else None)),
            (p.vendor_gstin if p else (g.vendor_gstin if g else None)),
            p.invoice_number if p else None,
            g.invoice_number if g else None,
            p.invoice_date.isoformat() if p else None,
            g.invoice_date.isoformat() if g else None,
            book_val,
            gstr_val,
            diff,
            r.plain_language_note,
            (exc.status.value if exc and hasattr(exc.status, "value")
             else (exc.status if exc else "—")),
            (exc.assigned_to if exc else None),
            (exc.resolution_notes if exc else None),
            (exc.approved_by if exc else None),
            (exc.itc_lapse_deadline.isoformat() if exc and exc.itc_lapse_deadline else None),
        ])
        status_val = r.match_status.value if hasattr(r.match_status, "value") else r.match_status
        color = STATUS_COLORS.get(status_val)
        if color:
            for cell in ws[row_idx]:
                cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        row_idx += 1

    widths = [22, 11, 26, 18, 20, 20, 16, 16, 16, 16, 14, 60, 18, 24, 40, 24, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
