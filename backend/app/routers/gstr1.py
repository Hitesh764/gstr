"""GSTR-1 preparation workspace: versioned uploads, review and draft export."""
from datetime import date, datetime
from calendar import monthrange
from collections import defaultdict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
import json
import re
import csv
from io import StringIO
from typing import Optional
from urllib.parse import quote, urlparse

import pandas as pd
from sqlalchemy import case, func, or_
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from app.matching.normalize import normalize_invoice_number
from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Request, UploadFile, Response
from pydantic import BaseModel, Field, field_validator
from fastapi.responses import StreamingResponse

from app import audit, models
from app.audit import AuditAction
from app.deps import Role, assert_client_visible, client_ip, get_tenant, require_exact_roles, require_role, user_agent
from app.ingestion.gstr1_sales import normalize_document_type, normalize_supply_type, parse_sales_register, reporting_category, validate_note_relationships
from app.ingestion.gstr1_annexure import build_annexure_report
from app.ingestion.gstr1_annexure_detail import build_hsn_summary, table_map_for_period, transaction_view
from app.storage import get_storage
from app.tenancy import TenantSession

router = APIRouter(prefix="/api/gstr1", tags=["GSTR-1 preparation"])
MAX_UPLOAD_BYTES = 15 * 1024 * 1024


def _known_sales_format_profiles(tenant: TenantSession, gstin_id: str, filename: str) -> list[dict]:
    source_type = "csv" if filename.casefold().endswith(".csv") else "excel"
    return [{"fingerprint": profile.fingerprint, "column_rules": profile.column_rules,
             "confirmed_columns": profile.confirmed_columns or {}}
            for profile in tenant.query(models.GSTR1ClientFormatProfile).filter_by(
                client_gstin_id=gstin_id, source_type=source_type).all()]


def _business_case(value):
    """Serialize workspace data with public business names, never ORM fields."""
    if isinstance(value, dict):
        return {re.sub(r"_([a-z])", lambda match: match.group(1).upper(), key): _business_case(item)
                for key, item in value.items()}
    if isinstance(value, list):
        return [_business_case(item) for item in value]
    return value


def _period(value: str) -> str:
    from app.ingestion.gstr1_sales import period_month
    try:
        period_month(value)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    return value


def _gst_rate_reference_rules(tenant: TenantSession) -> list[dict]:
    system_rules = [{"version": rule.version, "definition": rule.definition,
        "effective_from": rule.effective_from.isoformat(),
        "effective_to": rule.effective_to.isoformat() if rule.effective_to else None,
        "source_url": rule.source_url, "verified_at": rule.verified_at.isoformat() if rule.verified_at else None, "scope": "system"}
        for rule in tenant.db.query(models.GSTR1RuleDefinition)
        .filter(models.GSTR1RuleDefinition.rule_key == "GST_RATE_BY_HSN").all()]
    firm_rules = [{"version": rule.version, "definition": {"rates_by_hsn": rule.rates_by_hsn},
        "effective_from": rule.effective_from.isoformat(), "effective_to": rule.effective_to.isoformat() if rule.effective_to else None,
        "source_url": rule.source_url, "scope": "firm"}
        for rule in tenant.query(models.GSTR1TenantRateReference).filter_by(is_current=True).all()]
    return system_rules + firm_rules


def _gst_rate_reference_status(tenant: TenantSession, period: str, snapshot: dict | None = None) -> dict:
    from app.ingestion.gstr1_sales import period_month
    period_start = period_month(period)
    period_end = date(period_start.year, period_start.month, monthrange(period_start.year, period_start.month)[1])
    rules = _gst_rate_reference_rules(tenant)
    effective = [rule for rule in rules if rule["effective_from"] <= period_end.isoformat()
        and (not rule["effective_to"] or rule["effective_to"] >= period_start.isoformat())]
    firm_references = [{"version": rule.version, "source_url": rule.source_url,
                "effective_from": rule.effective_from.isoformat(), "effective_to": rule.effective_to.isoformat() if rule.effective_to else None,
                "hsn_count": len(rule.rates_by_hsn)} for rule in tenant.query(models.GSTR1TenantRateReference).filter_by(is_current=True).order_by(models.GSTR1TenantRateReference.effective_from.desc()).all()]
    if snapshot is None:
        detail = "This upload predates rate-reference snapshots. Re-import it to record which effective-dated reference was used."
        applicable = False
        versions = []
    else:
        detail = snapshot.get("detail", "Rate-reference status is unavailable.")
        applicable = bool(snapshot.get("applicable"))
        versions = snapshot.get("rule_versions", [])
        if not applicable and effective:
            detail = "A rate reference is now configured for this period, but it was not active when this register was imported. Re-import to run the comparison."
    return {"configured": bool(rules), "applicable": applicable, "rule_versions": versions,
        "firm_references": firm_references, "detail": detail}


def _invoice_tax_check(row: models.GSTR1SalesInvoice) -> dict | None:
    if row.gst_rate is None:
        return None
    expected = (row.taxable_value * row.gst_rate / Decimal("100")).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    reported = row.cgst + row.sgst + row.igst
    return {"enteredRate": float(row.gst_rate), "expectedGst": float(expected), "reportedGst": float(reported),
        "difference": float(abs(expected - reported))}


class GSTRateReferenceIn(BaseModel):
    version: str = Field(min_length=1, max_length=60)
    effective_from: date
    effective_to: date | None = None
    source_url: str = Field(min_length=8, max_length=500)
    rates_by_hsn: dict[str, object] = Field(min_length=1)

    @field_validator("version")
    @classmethod
    def clean_version(cls, value: str) -> str:
        value = value.strip()
        if not value:
            raise ValueError("Version is required.")
        return value

    @field_validator("source_url")
    @classmethod
    def valid_source_url(cls, value: str) -> str:
        value = value.strip()
        parsed = urlparse(value)
        if parsed.scheme not in {"http", "https"} or not parsed.netloc:
            raise ValueError("Source must be an http or https URL.")
        return value

    @field_validator("rates_by_hsn")
    @classmethod
    def valid_rate_map(cls, value: dict[str, object]) -> dict[str, list[str]]:
        result = {}
        for raw_code, raw_rates in value.items():
            code = re.sub(r"\s+", "", str(raw_code)).upper()
            if not re.fullmatch(r"[A-Z0-9]{4,8}", code):
                raise ValueError(f"HSN/SAC '{raw_code}' must contain 4 to 8 letters or digits.")
            options = raw_rates if isinstance(raw_rates, list) else [raw_rates]
            if not options:
                raise ValueError(f"At least one reference rate is required for HSN/SAC {code}.")
            rates = []
            for option in options:
                try:
                    rate = Decimal(str(option))
                except (InvalidOperation, ValueError):
                    raise ValueError(f"Rate '{option}' for HSN/SAC {code} is not numeric.")
                if not rate.is_finite() or not Decimal("0") <= rate <= Decimal("100"):
                    raise ValueError(f"Rate for HSN/SAC {code} must be between 0 and 100 percent.")
                rates.append(format(rate.normalize(), "f"))
            result[code] = sorted(set(rates), key=Decimal)
        if not result:
            raise ValueError("Add at least one HSN/SAC reference rate.")
        return result


@router.get("/rate-references")
def list_gst_rate_references(actor: models.User = Depends(require_role(Role.CLIENT_VIEWER)),
        tenant: TenantSession = Depends(get_tenant)):
    refs = tenant.query(models.GSTR1TenantRateReference).filter_by(is_current=True).order_by(
        models.GSTR1TenantRateReference.effective_from.desc()).all()
    return [{"id": row.id, "version": row.version, "source_url": row.source_url,
        "effective_from": row.effective_from.isoformat(), "effective_to": row.effective_to.isoformat() if row.effective_to else None,
        "hsn_count": len(row.rates_by_hsn), "uploaded_at": row.uploaded_at.isoformat()} for row in refs]


@router.post("/rate-references")
def create_gst_rate_reference(payload: GSTRateReferenceIn, request: Request,
        actor: models.User = Depends(require_role(Role.OWNER)), tenant: TenantSession = Depends(get_tenant)):
    if payload.effective_to and payload.effective_to < payload.effective_from:
        raise HTTPException(status_code=422, detail="Effective-to date must be on or after effective-from date.")
    if tenant.query(models.GSTR1TenantRateReference).filter_by(version=payload.version).first():
        raise HTTPException(status_code=409, detail="That reference version already exists. Use a new version name.")
    for prior in tenant.query(models.GSTR1TenantRateReference).filter_by(effective_from=payload.effective_from, is_current=True).all():
        prior.is_current = False
    ref = models.GSTR1TenantRateReference(id=models.gen_id(), version=payload.version,
        rates_by_hsn=payload.rates_by_hsn, source_url=payload.source_url,
        effective_from=payload.effective_from, effective_to=payload.effective_to, uploaded_by=actor.id)
    tenant.add(ref)
    tenant.flush()
    audit.record(tenant.db, firm_id=tenant.firm_id, user_id=actor.id,
        action=AuditAction.GSTR1_GST_RATE_REFERENCE_MAINTAINED, entity_type="GSTR1TenantRateReference", entity_id=ref.id,
        metadata={"version": ref.version, "effective_from": ref.effective_from.isoformat(),
            "effective_to": ref.effective_to.isoformat() if ref.effective_to else None,
            "hsn_count": len(ref.rates_by_hsn), "source_url": ref.source_url},
        ip_address=client_ip(request), user_agent=user_agent(request))
    tenant.commit()
    return {"id": ref.id, "version": ref.version, "effective_from": ref.effective_from.isoformat(),
        "effective_to": ref.effective_to.isoformat() if ref.effective_to else None,
        "hsn_count": len(ref.rates_by_hsn), "source_url": ref.source_url}


@router.post("/gstins/{gstin_id}/uploads")
async def upload_sales_register(gstin_id: str, request: Request,
        period: str = Query(..., description="YYYY-MM"), file: UploadFile = File(...),
        date_interpretation: Optional[str] = Form(None), date_column: Optional[str] = Form(None),
        column_choices: Optional[str] = Form(None),
        business_column_choices: Optional[str] = Form(None),
        actor: models.User = Depends(require_role(Role.ASSOCIATE)), tenant: TenantSession = Depends(get_tenant)):
    period = _period(period)
    gstin = tenant.get_or_404(models.ClientGSTIN, gstin_id, "GSTIN not found")
    name = file.filename or ""
    if not name.lower().endswith((".csv", ".xlsx", ".xls")):
        raise HTTPException(status_code=415, detail="Upload a CSV or Excel sales register.")
    content = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Sales register exceeds the 15 MB limit.")
    if not content:
        raise HTTPException(status_code=422, detail="Uploaded sales register is empty.")
    b2cl_rules = [{"version": rule.version, "definition": rule.definition,
                   "effective_from": rule.effective_from.isoformat(),
                   "effective_to": rule.effective_to.isoformat() if rule.effective_to else None}
                  for rule in tenant.db.query(models.GSTR1RuleDefinition)
                  .filter(models.GSTR1RuleDefinition.rule_key == "B2CL_THRESHOLD").all()]
    if not b2cl_rules:
        raise HTTPException(status_code=503, detail="Effective-dated GSTR-1 reference rules are not installed. Apply database migrations.")
    # Database/schema failures must remain server errors; don't report them as
    # malformed client files in the parser's business-friendly error handler.
    known_profiles = _known_sales_format_profiles(tenant, gstin_id, name)
    historical = (tenant.db.query(models.GSTR1SalesInvoice, models.GSTR1SalesUpload.period)
        .join(models.GSTR1SalesUpload, models.GSTR1SalesUpload.id == models.GSTR1SalesInvoice.upload_id)
        .filter(models.GSTR1SalesUpload.client_gstin_id == gstin_id, models.GSTR1SalesUpload.period < period,
            models.GSTR1SalesUpload.is_current.is_(True)).all())
    historical_note_rows = [{
        "customer_name": prior.customer_name, "customer_gstin": prior.customer_gstin,
        "invoice_number": prior.invoice_number, "normalized_invoice_number": prior.normalized_invoice_number,
        "invoice_date": prior.invoice_date, "invoice_value": prior.invoice_value,
        "taxable_value": prior.taxable_value, "cgst": prior.cgst, "sgst": prior.sgst,
        "igst": prior.igst, "cess": prior.cess, "document_type": prior.document_type,
        "invoice_status": prior.invoice_status, "original_invoice_number": prior.original_invoice_number,
        "period": source_period,
    } for prior, source_period in historical]
    try:
        rate_rules = _gst_rate_reference_rules(tenant)
        confirmed_columns = {}
        selected_business_columns = {}
        if date_interpretation is not None or date_column is not None:
            choices = {"invoiceDate": "invoice_date", "paymentDate": "payment_date",
                       "dueDate": "due_date", "otherDate": "other_date"}
            if date_interpretation not in choices or not date_column or len(date_column) > 150:
                raise HTTPException(status_code=422, detail="Choose what the selected date column represents, then try again.")
            confirmed_columns[date_column] = choices[date_interpretation]
        if column_choices:
            try:
                selected_columns = json.loads(column_choices)
            except (TypeError, json.JSONDecodeError) as exc:
                raise HTTPException(status_code=422, detail="Please choose the matching sales information from the file headings.") from exc
            field_choices = {"invoiceNumber": "invoice_number", "invoiceDate": "invoice_date", "taxableAmount": "taxable_value"}
            if not isinstance(selected_columns, dict) or any(key not in field_choices or not isinstance(value, str) or not value or len(value) > 150
                    for key, value in selected_columns.items()):
                raise HTTPException(status_code=422, detail="Please choose the matching sales information from the file headings.")
            for public_field, header in selected_columns.items():
                if public_field in {"invoiceNumber", "invoiceDate", "taxableAmount"}:
                    confirmed_columns[header] = field_choices[public_field]
        if business_column_choices:
            try:
                selected_business_columns = json.loads(business_column_choices)
            except (TypeError, json.JSONDecodeError) as exc:
                raise HTTPException(status_code=422, detail="Please choose one of the columns shown in the conflict review.") from exc
            business_field_choices = {"invoiceNumber": "invoice_number", "invoiceDate": "invoice_date", "taxableAmount": "taxable_value",
                "customerName": "customer_name", "customerGstin": "customer_gstin", "placeOfSupply": "place_of_supply",
                "cgst": "cgst", "sgst": "sgst", "igst": "igst", "cess": "cess", "invoiceValue": "invoice_value",
                "hsnSac": "hsn_sac", "gstRate": "gst_rate", "documentType": "document_type", "supplyType": "supply_type"}
            if not isinstance(selected_business_columns, dict) or any(key not in business_field_choices or
                    not isinstance(value, str) or not value or len(value) > 150
                    for key, value in selected_business_columns.items()):
                raise HTTPException(status_code=422, detail="Please choose one of the columns shown in the conflict review.")
            selected_business_columns = {header: business_field_choices[field]
                for field, header in selected_business_columns.items()}
        parsed = parse_sales_register(content, name, period, gstin.gstin, b2cl_rules, rate_rules,
            known_profiles=known_profiles, confirmed_columns=confirmed_columns,
            historical_invoices=historical_note_rows, business_column_choices=selected_business_columns)
    except HTTPException:
        raise
    except ValueError as exc:
        from app.ingestion.column_resolution import AmbiguousColumnError, UnresolvedColumnsError
        from app.ingestion.column_resolution import AmbiguousBusinessColumnError
        if isinstance(exc, AmbiguousBusinessColumnError):
            public_fields = {"customer_name": "customerName", "customer_gstin": "customerGstin",
                "invoice_number": "invoiceNumber", "invoice_date": "invoiceDate", "taxable_value": "taxableAmount",
                "cgst": "cgst", "sgst": "sgst", "igst": "igst", "cess": "cess", "invoice_value": "invoiceValue",
                "place_of_supply": "placeOfSupply", "hsn_sac": "hsnSac", "gst_rate": "gstRate",
                "document_type": "documentType", "supply_type": "supplyType"}
            raise HTTPException(status_code=422, detail={"message": str(exc), "businessConflict": {
                "field": public_fields.get(exc.field, "salesDetail"), "label": exc.field_label,
                "columns": exc.columns, "affectedRowCount": exc.affected_row_count,
                "sampleRows": exc.sample_rows}}) from exc
        if isinstance(exc, AmbiguousColumnError):
            raise HTTPException(status_code=422, detail={"message": "We found a date column but need to know what date it contains.",
                "ambiguousDate": {"column": exc.column, "options": [
                    {"value": "invoiceDate", "label": "Invoice date"},
                    {"value": "paymentDate", "label": "Payment date"},
                    {"value": "dueDate", "label": "Due date"},
                    {"value": "otherDate", "label": "Another date"}]}}) from exc
        if isinstance(exc, UnresolvedColumnsError):
            labels = {"invoiceNumber": "Invoice number", "invoiceDate": "Invoice date", "taxableAmount": "Taxable amount"}
            raise HTTPException(status_code=422, detail={"message": str(exc), "unresolvedFields": [
                {"value": field, "label": labels[field], "columns": exc.available_columns} for field in exc.fields]}) from exc
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except Exception as exc:
        # Keep parser/library details in server logs, not in the accountant workflow.
        import logging
        logging.getLogger(__name__).exception("Could not process GSTR-1 sales register")
        raise HTTPException(status_code=422,
            detail="We couldn't read this file. Check that it is a valid CSV or Excel sales register, then try again.") from exc
    format_profile = parsed.pop("_format_profile")
    historical_gstins_by_name = {}
    historical_hsn_by_item = {}
    for prior, _source_period in historical:
        name_key = re.sub(r"[^a-z0-9]+", "", (prior.customer_name or "").casefold())
        if name_key and prior.customer_gstin:
            historical_gstins_by_name.setdefault(name_key, set()).add(prior.customer_gstin)
        item_key = (name_key, re.sub(r"[^a-z0-9]+", "", (prior.description or "").casefold()))
        if item_key[0] and item_key[1] and prior.hsn_sac:
            historical_hsn_by_item.setdefault(item_key, set()).add(prior.hsn_sac)
    for row in parsed["rows"]:
        name_key = re.sub(r"[^a-z0-9]+", "", (row.get("customer_name") or "").casefold())
        old_gstins = historical_gstins_by_name.get(name_key, set()) if name_key else set()
        if old_gstins and row.get("customer_gstin") and row["customer_gstin"] not in old_gstins:
            issue = {"row": row["row_number"], "code": "CUSTOMER_GSTIN_DIFFERS_FROM_HISTORY", "severity": "REVIEW",
                "message": "This customer name has a different GSTIN in an earlier period. Verify the customer master; registrations can change."}
            parsed["issues"].append(issue)
            row["issues"].append({k: v for k, v in issue.items() if k != "row"})
            row["status"] = "SYSTEM_FLAGGED"
        item_key = (name_key, re.sub(r"[^a-z0-9]+", "", (row.get("description") or "").casefold()))
        old_hsns = historical_hsn_by_item.get(item_key, set())
        if old_hsns and row.get("hsn_sac") and row["hsn_sac"] not in old_hsns:
            issue = {"row": row["row_number"], "code": "CUSTOMER_HSN_CHANGED_FROM_HISTORY", "severity": "REVIEW",
                "message": "This customer and description used a different HSN/SAC in an earlier period. Check product classification and supporting facts."}
            parsed["issues"].append(issue)
            row["issues"].append({k: v for k, v in issue.items() if k != "row"})
            row["status"] = "SYSTEM_FLAGGED"
    version = (tenant.query(models.GSTR1SalesUpload)
        .filter(models.GSTR1SalesUpload.client_gstin_id == gstin_id, models.GSTR1SalesUpload.period == period)
        .order_by(models.GSTR1SalesUpload.version.desc()).first())
    next_version = (version.version + 1) if version else 1
    if version:
        version.is_current = False
    try:
        source_key = get_storage().save(firm_id=tenant.firm_id, category="gstr1_sales", data=content)
    except Exception as exc:
        raise HTTPException(status_code=500, detail="Could not securely store the uploaded sales register.") from exc
    profile = (tenant.query(models.GSTR1ClientFormatProfile).filter_by(client_gstin_id=gstin_id,
        source_type=format_profile["source_type"], fingerprint=format_profile["fingerprint"]).first())
    if profile is None:
        profile = models.GSTR1ClientFormatProfile(id=models.gen_id(), client_gstin_id=gstin_id, source_type=format_profile["source_type"],
            source_name="Sales Register", fingerprint=format_profile["fingerprint"], source_columns=format_profile["source_columns"],
            column_rules=format_profile["column_rules"], confirmed_columns=format_profile["confirmed_columns"],
            alias_version=format_profile["alias_version"], version=1,
            import_count=0, success_count=0)
        tenant.add(profile)
    else:
        mapping_changed = (profile.alias_version != format_profile["alias_version"] or
            profile.column_rules != format_profile["column_rules"] or
            any((profile.confirmed_columns or {}).get(header) != field
                for header, field in format_profile["confirmed_columns"].items()))
        profile.source_columns = format_profile["source_columns"]
        profile.column_rules = format_profile["column_rules"]
        profile.confirmed_columns = {**(profile.confirmed_columns or {}), **format_profile["confirmed_columns"]}
        profile.alias_version = format_profile["alias_version"]
        if mapping_changed:
            profile.version += 1
    profile.import_count += 1
    profile.success_count += 1
    profile.last_used_at = datetime.utcnow()
    normalization_snapshot = {**format_profile, "profile_id": profile.id, "profile_version": profile.version}
    upload = models.GSTR1SalesUpload(client_gstin_id=gstin_id, period=period, version=next_version,
        is_current=True, original_filename=name, file_path=source_key, source_sha256=parsed["source_sha256"],
        row_count=len(parsed["rows"]), issue_count=len(parsed["issues"]), uploaded_by=actor.id,
        rate_reference_snapshot=parsed["rate_reference_snapshot"], normalization_snapshot=normalization_snapshot)
    tenant.add(upload)
    tenant.flush()
    ret = (tenant.query(models.GSTR1Return).filter_by(client_gstin_id=gstin_id, period=period).first())
    if ret is None:
        ret = models.GSTR1Return(client_gstin_id=gstin_id, period=period, created_by=actor.id)
        tenant.add(ret)
    # Re-imports always return the work to draft, preserving old upload versions.
    ret.status = models.GSTR1ReturnStatus.DRAFT
    ret.created_by = actor.id
    ret.reviewed_by = None
    ret.reviewed_at = None
    ret.approved_by = None
    ret.approved_at = None
    ret.source_upload_id = upload.id
    ret.rule_version = parsed["rule_version"]
    ret.reconciliation_result = None
    # A changed invoice population makes prior Table 13 cross-checks stale.
    for document_record in tenant.query(models.GSTR1DocumentIssued).filter_by(return_id=ret.id, is_current=True).all():
        document_record.is_current = False
    tenant.flush()
    invoices_by_row = {}
    for index, item in enumerate(parsed["rows"], start=1):
        item.pop("issues")
        item.pop("linked_original_invoice", None)
        item.pop("note_relationship_status", None)
        invoice = models.GSTR1SalesInvoice(id=models.gen_id(), upload_id=upload.id, **item)
        tenant.add(invoice)
        invoices_by_row[invoice.row_number] = invoice
        if index % 500 == 0:
            tenant.flush()
    for issue in parsed["issues"]:
        row_number = issue.get("row")
        invoice = invoices_by_row.get(row_number)
        tenant.add(models.GSTR1ValidationIssue(upload_id=upload.id, invoice_id=invoice.id if invoice else None,
            row_number=row_number, code=issue["code"], severity=issue["severity"], message=issue["message"]))
    tenant.flush()
    audit.record(tenant.db, firm_id=tenant.firm_id, user_id=actor.id,
        action=AuditAction.GSTR1_SALES_REGISTER_UPLOADED, entity_type="GSTR1SalesUpload", entity_id=upload.id,
        metadata={"gstin": gstin.gstin, "period": period, "version": next_version,
                  "rows": len(parsed["rows"]), "issues": len(parsed["issues"]), "sha256": parsed["source_sha256"],
                  "column_resolution": format_profile["resolution_decisions"]},
        ip_address=client_ip(request), user_agent=user_agent(request))
    try:
        tenant.commit()
    except Exception:
        tenant.rollback()
        try:
            get_storage().delete(source_key, firm_id=tenant.firm_id)
        except Exception:
            pass
        raise
    flagged_rows = {issue.get("row") for issue in parsed["issues"] if issue.get("severity") in {"REVIEW", "ERROR"}}
    return {"uploadId": upload.id, "returnId": ret.id, "period": period, "version": next_version,
            "status": "DRAFT",
            "recordsProcessed": len(parsed["rows"]),
            "recordsAccepted": max(0, len(parsed["rows"]) - len(flagged_rows.intersection({row["row_number"] for row in parsed["rows"]}))),
            "recordsRequiringReview": len(flagged_rows),
            "criticalIssueCount": sum(issue.get("severity") == "ERROR" for issue in parsed["issues"]),
            "message": f"{len(parsed['rows']):,} sales records processed. {len(flagged_rows):,} need attention."}


@router.get("/gstins/{gstin_id}/returns/{period}")
def get_return(gstin_id: str, period: str, actor: models.User = Depends(require_role(Role.CLIENT_VIEWER)),
               tenant: TenantSession = Depends(get_tenant)):
    return _business_case(_get_return_data(gstin_id, period, actor, tenant))


@router.get("/gstins/{gstin_id}/returns/{period}/annexure")
def get_annexure_data(gstin_id: str, period: str, request: Request,
        report_type: str = Query("sales", pattern="^(sales|b2b|b2c|exports|notes|hsn)$"),
        page: int = Query(1, ge=1), page_size: int = Query(50, ge=1, le=100),
        search: str | None = Query(None, max_length=120), customer_gstin: str | None = Query(None, max_length=15),
        invoice_number: str | None = Query(None, max_length=100), place_of_supply: str | None = Query(None, max_length=2),
        category: str | None = Query(None, max_length=30),
        issue_code: str | None = Query(None, max_length=60), status: str | None = Query(None, pattern="^(READY|REVIEW_REQUIRED|ERROR)$"),
        gst_rate: Decimal | None = Query(None, ge=0, le=100), date_from: date | None = Query(None), date_to: date | None = Query(None),
        actor: models.User = Depends(require_role(Role.CLIENT_VIEWER)), tenant: TenantSession = Depends(get_tenant)):
    """Paginated annexure built exclusively from normalized imported transactions."""
    period = _period(period)
    gstin = tenant.get_or_404(models.ClientGSTIN, gstin_id, "GSTIN not found")
    assert_client_visible(actor, gstin.client_id)
    ret = tenant.query(models.GSTR1Return).filter_by(client_gstin_id=gstin_id, period=period).first()
    if not ret:
        raise HTTPException(status_code=404, detail="No GSTR-1 draft exists for this period.")
    upload = tenant.get(models.GSTR1SalesUpload, ret.source_upload_id) if ret.source_upload_id else None
    if not upload:
        return _business_case({"report_type": report_type, "period": period, "rules_version": ret.rule_version,
            "table_mapping_version": table_map_for_period(period)["version"],
            "summary": {"source_transactions": 0, "processed_transactions": 0, "excluded_transactions": 0,
                "valid_transactions": 0, "warnings": 0, "errors": 0, "duplicates": 0, "requires_review": 0,
                "taxable_value": 0, "cgst": 0, "sgst": 0, "igst": 0, "cess": 0},
            "pagination": {"page": page, "page_size": page_size, "total": 0, "pages": 0}, "rows": []})

    base = tenant.query(models.GSTR1SalesInvoice).filter_by(upload_id=upload.id)
    Invoice = models.GSTR1SalesInvoice
    Issue = models.GSTR1ValidationIssue
    if report_type == "b2b":
        base = base.filter(Invoice.classification == "B2B")
    elif report_type == "b2c":
        base = base.filter(Invoice.classification.in_(["B2CL", "B2CS"]))
    elif report_type == "exports":
        base = base.filter(Invoice.classification.in_(["EXPORT", "SEZ", "DEEMED_EXPORT"]))
    elif report_type == "notes":
        base = base.filter(or_(Invoice.document_type.like("CREDIT_NOTE%"), Invoice.document_type.like("DEBIT_NOTE%")))
    if search:
        term = f"%{search.strip()}%"
        base = base.filter(or_(Invoice.invoice_number.ilike(term), Invoice.customer_name.ilike(term),
            Invoice.customer_gstin.ilike(term), Invoice.hsn_sac.ilike(term), Invoice.port_code.ilike(term),
            Invoice.shipping_bill_number.ilike(term)))
    if customer_gstin:
        base = base.filter(Invoice.customer_gstin.ilike(f"%{customer_gstin.strip()}%"))
    if invoice_number:
        base = base.filter(Invoice.invoice_number.ilike(f"%{invoice_number.strip()}%"))
    if place_of_supply:
        base = base.filter(Invoice.place_of_supply == place_of_supply.strip())
    if category:
        base = base.filter(Invoice.classification == category.upper())
    if issue_code:
        base = base.filter(tenant.query(Issue).filter(Issue.invoice_id == Invoice.id,
            Issue.code == issue_code.upper(), Issue.status == "OPEN").exists())
    if gst_rate is not None:
        base = base.filter(Invoice.gst_rate == gst_rate)
    if date_from:
        base = base.filter(Invoice.invoice_date >= date_from)
    if date_to:
        base = base.filter(Invoice.invoice_date <= date_to)
    open_check = tenant.query(Issue).filter(Issue.invoice_id == Invoice.id, Issue.status == "OPEN").exists()
    error_check = tenant.query(Issue).filter(Issue.invoice_id == Invoice.id, Issue.status == "OPEN", Issue.severity == "ERROR").exists()
    if status == "ERROR":
        base = base.filter(error_check)
    elif status == "REVIEW_REQUIRED":
        base = base.filter(open_check, ~error_check)
    elif status == "READY":
        base = base.filter(~open_check)

    total = base.count()
    seller_gstin = gstin.gstin
    b2cl_rules = [{"version": rule.version, "definition": rule.definition,
        "effective_from": rule.effective_from.isoformat(), "effective_to": rule.effective_to.isoformat() if rule.effective_to else None}
        for rule in tenant.db.query(models.GSTR1RuleDefinition).filter(models.GSTR1RuleDefinition.rule_key == "B2CL_THRESHOLD").all()]

    # Summary figures cover the full imported register, independent of the active tab/filter.
    all_query = tenant.query(Invoice).filter_by(upload_id=upload.id)
    note_sign = case((Invoice.document_type.like("CREDIT_NOTE%"), -1), else_=1)
    aggregates = all_query.with_entities(func.count(Invoice.id), func.coalesce(func.sum(Invoice.taxable_value * note_sign), 0),
        func.coalesce(func.sum(Invoice.cgst * note_sign), 0), func.coalesce(func.sum(Invoice.sgst * note_sign), 0),
        func.coalesce(func.sum(Invoice.igst * note_sign), 0), func.coalesce(func.sum(Invoice.cess * note_sign), 0)).one()
    issue_counts = tenant.query(Issue).with_entities(Issue.severity, func.count(Issue.id)).filter_by(upload_id=upload.id, status="OPEN").group_by(Issue.severity).all()
    open_errors = sum(count for severity, count in issue_counts if severity == "ERROR")
    open_warnings = sum(count for severity, count in issue_counts if severity != "ERROR")
    excluded = tenant.query(Issue).with_entities(func.count(func.distinct(Issue.row_number))).filter_by(upload_id=upload.id, invoice_id=None).scalar() or 0
    duplicates = tenant.query(Issue).with_entities(func.count(func.distinct(Issue.invoice_id))).filter(Issue.upload_id == upload.id,
        Issue.invoice_id.isnot(None), Issue.status == "OPEN", Issue.code.in_(["POSSIBLE_DUPLICATE", "POSSIBLE_FUZZY_DUPLICATE", "POSSIBLE_DUPLICATE_NOTE", "INVOICE_NUMBER_REUSED"])).scalar() or 0
    needs_review = tenant.query(Issue).with_entities(func.count(func.distinct(Issue.invoice_id))).filter(Issue.upload_id == upload.id,
        Issue.invoice_id.isnot(None), Issue.status == "OPEN").scalar() or 0
    ready = max(0, int(aggregates[0]) - needs_review)
    category_counts = {key: count for key, count in all_query.with_entities(Invoice.classification, func.count(Invoice.id)).group_by(Invoice.classification).all()}
    summary = {"source_transactions": int(aggregates[0]) + excluded, "processed_transactions": int(aggregates[0]),
        "excluded_transactions": excluded, "valid_transactions": ready, "warnings": open_warnings,
        "errors": open_errors, "duplicates": duplicates, "requires_review": needs_review,
        "b2b": category_counts.get("B2B", 0), "b2c_large": category_counts.get("B2CL", 0),
        "b2c_others": category_counts.get("B2CS", 0), "exports": category_counts.get("EXPORT", 0),
        "sez": category_counts.get("SEZ", 0), "credit_notes": category_counts.get("CREDIT_NOTE", 0),
        "debit_notes": category_counts.get("DEBIT_NOTE", 0), "taxable_value": float(aggregates[1]),
        "cgst": float(aggregates[2]), "sgst": float(aggregates[3]), "igst": float(aggregates[4]), "cess": float(aggregates[5])}

    if report_type == "hsn":
        grouped = base.with_entities(Invoice.hsn_sac, Invoice.description, Invoice.uqc, Invoice.gst_rate, Invoice.classification,
            func.count(func.distinct(Invoice.invoice_number)), func.count(Invoice.id), func.sum(Invoice.quantity * note_sign), func.sum(Invoice.taxable_value * note_sign),
            func.sum(Invoice.cgst * note_sign), func.sum(Invoice.sgst * note_sign),
            func.sum(Invoice.igst * note_sign), func.sum(Invoice.cess * note_sign)).group_by(
                Invoice.hsn_sac, Invoice.description, Invoice.gst_rate, Invoice.uqc, Invoice.classification)
        total = grouped.count()
        rows = [{"hsn_sac": item[0] or "Missing", "description": item[1], "uqc": item[2], "gst_rate": float(item[3]) if item[3] is not None else None,
            "reporting_group": item[4], "documents": int(item[5]), "records": int(item[6]),
            "quantity": float(item[7] or 0), "taxable_value": float(item[8] or 0), "cgst": float(item[9] or 0),
            "sgst": float(item[10] or 0), "igst": float(item[11] or 0), "cess": float(item[12] or 0)}
            for item in grouped.order_by(Invoice.hsn_sac, Invoice.classification, Invoice.gst_rate).offset((page - 1) * page_size).limit(page_size).all()]
    else:
        selected = base.order_by(Invoice.invoice_date, Invoice.invoice_number, Invoice.row_number).offset((page - 1) * page_size).limit(page_size).all()
        selected_ids = [row.id for row in selected]
        issue_rows = tenant.query(Issue).filter(Issue.invoice_id.in_(selected_ids)).order_by(Issue.id).all() if selected_ids else []
        issues_by_invoice = defaultdict(list)
        for issue in issue_rows:
            issues_by_invoice[issue.invoice_id].append(issue)
        rows = [transaction_view(row, issues_by_invoice[row.id], seller_gstin, period, b2cl_rules) for row in selected]
    return _business_case({"report_type": report_type, "period": period, "rules_version": ret.rule_version,
        "table_mapping_version": table_map_for_period(period)["version"], "summary": summary,
        "pagination": {"page": page, "page_size": page_size, "total": total, "pages": (total + page_size - 1) // page_size},
        "rows": rows})


def _get_return_data(gstin_id: str, period: str, actor: models.User, tenant: TenantSession):
    period = _period(period)
    gstin = tenant.get_or_404(models.ClientGSTIN, gstin_id, "GSTIN not found")
    assert_client_visible(actor, gstin.client_id)
    ret = tenant.query(models.GSTR1Return).filter_by(client_gstin_id=gstin_id, period=period).first()
    if not ret:
        raise HTTPException(status_code=404, detail="No GSTR-1 draft exists for this period.")
    upload = tenant.get(models.GSTR1SalesUpload, ret.source_upload_id) if ret.source_upload_id else None
    invoices = tenant.query(models.GSTR1SalesInvoice).filter_by(upload_id=upload.id).all() if upload else []
    issue_rows = tenant.query(models.GSTR1ValidationIssue).filter_by(upload_id=upload.id).all() if upload else []
    document_records = tenant.query(models.GSTR1DocumentIssued).filter_by(return_id=ret.id, is_current=True).all()
    rate_reference = _gst_rate_reference_status(tenant, period, upload.rate_reference_snapshot if upload else None)
    issue_by_invoice = {}
    for issue in issue_rows:
        issue_by_invoice.setdefault(issue.invoice_id, []).append(issue)
    seller_gstin = tenant.get(models.ClientGSTIN, gstin_id).gstin
    b2cl_rules = [{"version": rule.version, "definition": rule.definition,
        "effective_from": rule.effective_from.isoformat(),
        "effective_to": rule.effective_to.isoformat() if rule.effective_to else None}
        for rule in tenant.db.query(models.GSTR1RuleDefinition)
        .filter(models.GSTR1RuleDefinition.rule_key == "B2CL_THRESHOLD").all()]
    def row_reporting_category(row, tax_period=period):
        return reporting_category(row.document_type, row.supply_type, row.classification,
            seller_gstin=seller_gstin, customer_gstin=row.customer_gstin, place_of_supply=row.place_of_supply,
            invoice_value=row.invoice_value, period=tax_period, b2cl_rules=b2cl_rules)
    by_category = {}
    for row in invoices:
        category = row_reporting_category(row)
        agg = by_category.setdefault(category, {"count": 0, "taxable_value": 0.0, "tax": 0.0})
        agg["count"] += 1
        agg["taxable_value"] += float(row.taxable_value)
        agg["tax"] += float(row.cgst + row.sgst + row.igst + row.cess)
    document_type_summary = {}
    for row in invoices:
        doc_type = normalize_document_type(row.document_type)
        bucket = document_type_summary.setdefault(doc_type, {"count": 0, "taxable_value": 0.0, "tax": 0.0})
        bucket["count"] += 1
        bucket["taxable_value"] += float(row.taxable_value)
        bucket["tax"] += float(row.cgst + row.sgst + row.igst + row.cess)
    note_relationship_rows = [{"row_number": row.row_number, "customer_name": row.customer_name,
        "customer_gstin": row.customer_gstin, "invoice_number": row.invoice_number,
        "normalized_invoice_number": row.normalized_invoice_number, "invoice_date": row.invoice_date,
        "invoice_value": row.invoice_value, "taxable_value": row.taxable_value, "cgst": row.cgst,
        "sgst": row.sgst, "igst": row.igst, "cess": row.cess, "document_type": row.document_type,
        "invoice_status": row.invoice_status, "original_invoice_number": row.original_invoice_number,
        "original_invoice_date": row.original_invoice_date, "issues": [], "status": "NEEDS_REVIEW"}
        for row in invoices]
    note_relationship_history = []
    if any(normalize_document_type(row.document_type).startswith(("CREDIT_NOTE", "DEBIT_NOTE")) for row in invoices):
        prior_records = (tenant.db.query(models.GSTR1SalesInvoice, models.GSTR1SalesUpload.period)
            .join(models.GSTR1SalesUpload, models.GSTR1SalesUpload.id == models.GSTR1SalesInvoice.upload_id)
            .filter(models.GSTR1SalesUpload.client_gstin_id == gstin_id,
                models.GSTR1SalesUpload.period < period, models.GSTR1SalesUpload.is_current.is_(True)).all())
        note_relationship_history = [{
            "customer_name": prior.customer_name, "customer_gstin": prior.customer_gstin,
            "invoice_number": prior.invoice_number, "normalized_invoice_number": prior.normalized_invoice_number,
            "invoice_date": prior.invoice_date, "invoice_value": prior.invoice_value,
            "taxable_value": prior.taxable_value, "cgst": prior.cgst, "sgst": prior.sgst,
            "igst": prior.igst, "cess": prior.cess, "document_type": prior.document_type,
            "invoice_status": prior.invoice_status, "original_invoice_number": prior.original_invoice_number,
            "period": source_period,
        } for prior, source_period in prior_records]
    note_relationships = validate_note_relationships(note_relationship_rows, [], note_relationship_history, period)
    hsn_categories = {"B2B": "B2B", "B2CL": "B2C Large", "B2CS": "B2C Others", "EXPORT": "Export",
        "SEZ": "SEZ", "DEEMED_EXPORT": "Deemed export", "EXEMPT": "Nil/exempt", "OTHER": "Needs classification",
        "UNCLASSIFIED": "Needs classification"}
    hsn_summary = {}
    hsn_rows = build_hsn_summary([{
        "id": row.id, "hsnSac": row.hsn_sac, "description": row.description, "uqc": row.uqc,
        "gstRate": float(row.gst_rate) if row.gst_rate is not None else None,
        "reportingGroup": hsn_categories.get(row_reporting_category(row), row_reporting_category(row)),
        "invoiceNumber": row.invoice_number, "customerGstin": row.customer_gstin,
        "documentType": "Credit note" if normalize_document_type(row.document_type).startswith("CREDIT_NOTE") else
            "Debit note" if normalize_document_type(row.document_type).startswith("DEBIT_NOTE") else "Invoice",
        "quantity": float(row.quantity or 0), "taxableValue": float(row.taxable_value),
        "cgst": float(row.cgst), "sgst": float(row.sgst), "igst": float(row.igst), "cess": float(row.cess),
    } for row in invoices])
    for item in hsn_rows:
        legacy_group = "B2C" if item["reportingGroup"].startswith("B2C") else item["reportingGroup"]
        key = f"{item['hsnSac']}|{legacy_group}"
        if key in hsn_summary:
            # Preserve the historic key where possible, while keeping distinct
            # rates, descriptions, units and B2C categories from overwriting one another.
            key += "|" + "|".join(str(value or "") for value in
                (item["reportingGroup"], item["gstRate"], item["uqc"], item["description"]))
        hsn_summary[key] = {
            "hsn_sac": item["hsnSac"], "description": item["description"], "uqc": item["uqc"],
            "gst_rate": item["gstRate"], "reporting_group": item["reportingGroup"],
            "invoice_count": item["documents"], "record_count": item["records"], "quantity": item["quantity"],
            "taxable_value": item["taxableValue"], "cgst": item["cgst"], "sgst": item["sgst"],
            "igst": item["igst"], "cess": item["cess"],
        }
    prior_returns = (tenant.query(models.GSTR1Return)
        .filter(models.GSTR1Return.client_gstin_id == gstin_id, models.GSTR1Return.period < period,
            models.GSTR1Return.source_upload_id.isnot(None))
        .order_by(models.GSTR1Return.period.desc()).limit(6).all())
    trend = []
    for old_return in reversed(prior_returns):
        old_rows = tenant.query(models.GSTR1SalesInvoice).filter_by(upload_id=old_return.source_upload_id).all()
        old_taxable = sum((row.taxable_value for row in old_rows), Decimal("0"))
        old_b2b = sum(1 for row in old_rows if row_reporting_category(row, old_return.period) == "B2B")
        trend.append({"period": old_return.period, "invoice_count": len(old_rows), "taxable_value": float(old_taxable),
            "b2b_share": old_b2b / len(old_rows) if old_rows else 0})
    current_taxable = sum((row.taxable_value for row in invoices), Decimal("0"))
    anomaly_hints = []
    if trend and current_taxable and sum(x["taxable_value"] for x in trend) > 0:
        average_taxable = sum(x["taxable_value"] for x in trend) / len(trend)
        if current_taxable > average_taxable * 1.75 or current_taxable < average_taxable * 0.5:
            anomaly_hints.append("Current taxable sales differ materially from the average of available prior periods.")
        current_b2b_share = sum(1 for row in invoices if row_reporting_category(row) == "B2B") / len(invoices) if invoices else 0
        average_b2b_share = sum(x["b2b_share"] for x in trend) / len(trend)
        if abs(current_b2b_share - average_b2b_share) >= 0.25:
            anomaly_hints.append("The B2B share differs materially from the average of available prior periods.")
    annexure_rows = [{"document_type": normalize_document_type(row.document_type),
        "supply_type": normalize_supply_type(row.supply_type), "category": row_reporting_category(row),
        "customer_gstin": row.customer_gstin, "taxable_value": row.taxable_value,
        "cgst": row.cgst, "sgst": row.sgst, "igst": row.igst, "cess": row.cess}
        for row in invoices]
    annexure_period_label = date.fromisoformat(f"{period}-01").strftime("%b-%y") + " turnover"
    annexure_report = build_annexure_report(annexure_rows, annexure_period_label)
    return {"id": ret.id, "period": ret.period, "status": ret.status.value, "rule_version": ret.rule_version,
        "reconciliation": ret.reconciliation_result, "gst_rate_reference": rate_reference,
        "historical_trend": {"periods": trend, "anomaly_hints": anomaly_hints, "basis": "Up to six previously imported periods; review prompt only."},
        "reviewed_by": ret.reviewed_by, "reviewed_at": ret.reviewed_at.isoformat() if ret.reviewed_at else None,
        "approved_by": ret.approved_by, "approved_at": ret.approved_at.isoformat() if ret.approved_at else None,
        "upload": {"id": upload.id, "version": upload.version, "filename": upload.original_filename,
                   "row_count": upload.row_count, "sha256": upload.source_sha256} if upload else None,
        "summary": {"invoice_count": len(invoices), "issue_count": len(issue_rows), "categories": by_category,
                    "document_types": document_type_summary,
                    "open_issues": sum(1 for issue in issue_rows if issue.status == "OPEN")},
        "annexure_report": annexure_report,
        "upload_issues": [{"row_number": x.row_number, "code": x.code, "severity": x.severity,
            "message": x.message, "status": x.status} for x in issue_rows if x.invoice_id is None],
        "documents_issued": [{"id": r.id, "document_type": r.document_type, "series_from": r.series_from,
            "series_to": r.series_to, "issued_count": r.issued_count, "cancelled_count": r.cancelled_count,
            "net_issued": r.issued_count-r.cancelled_count, "version": r.version,
            "prepared_by": r.prepared_by, "reviewed_by": r.reviewed_by,
            "reviewed_at": r.reviewed_at.isoformat() if r.reviewed_at else None} for r in document_records],
        "hsn_summary": hsn_summary,
        "invoices": [{"id": row.id, "sourceRowNumber": row.row_number, "invoiceNumber": row.invoice_number,
            "invoiceDate": row.invoice_date.isoformat(), "customer": row.customer_name,
            "customerGstin": row.customer_gstin, "invoiceValue": float(row.invoice_value),
            "taxableValue": float(row.taxable_value), "cgstAmount": float(row.cgst), "sgstAmount": float(row.sgst),
            "igstAmount": float(row.igst), "cessAmount": float(row.cess), "gstRate": float(row.gst_rate) if row.gst_rate is not None else None,
            "taxCheck": _invoice_tax_check(row), "hsnSac": row.hsn_sac,
            "description": row.description, "originalInvoiceNumber": row.original_invoice_number,
            "invoiceStatus": row.invoice_status, "placeOfSupply": row.place_of_supply,
            "documentType": normalize_document_type(row.document_type),
            "supplyType": normalize_supply_type(row.supply_type),
            "category": row_reporting_category(row), "categoryReason": (
                f"{row_reporting_category(row)} classification inferred from recipient and supply details; source document type needs review."
                if normalize_document_type(row.document_type) == "OTHER" else row.classification_reason),
            "noteRelationship": note_relationships.get(row.row_number),
            "reviewStatus": row.status.value, "issues": [{"id": x.id, "code": x.code, "severity": x.severity,
                "message": x.message, "status": x.status, "review_decision": x.review_decision,
            "review_note": x.review_note} for x in issue_by_invoice.get(row.id, [])],
            "reviewed_by": row.reviewed_by, "reviewed_at": row.reviewed_at.isoformat() if row.reviewed_at else None} for row in invoices]}


@router.post("/returns/{return_id}/reconcile-books")
async def reconcile_books_register(return_id: str, request: Request,
        file: UploadFile = File(...), actor: models.User = Depends(require_role(Role.ASSOCIATE)),
        tenant: TenantSession = Depends(get_tenant)):
    """Compare a separately supplied books sales register with the current GSTR-1 preparation."""
    ret = tenant.get_or_404(models.GSTR1Return, return_id, "GSTR-1 return not found")
    if ret.status != models.GSTR1ReturnStatus.DRAFT:
        raise HTTPException(status_code=409, detail="Books reconciliation can only be updated while the return is a draft.")
    gstin = tenant.get_or_404(models.ClientGSTIN, ret.client_gstin_id, "GSTIN not found")
    assert_client_visible(actor, gstin.client_id)
    name = file.filename or ""
    if not name.lower().endswith((".csv", ".xlsx", ".xls")):
        raise HTTPException(status_code=415, detail="Upload the separate books sales register as CSV or Excel.")
    content = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Books register exceeds the 15 MB limit.")
    rules = [{"version": r.version, "definition": r.definition, "effective_from": r.effective_from.isoformat(),
        "effective_to": r.effective_to.isoformat() if r.effective_to else None}
        for r in tenant.db.query(models.GSTR1RuleDefinition).filter_by(rule_key="B2CL_THRESHOLD").all()]
    try:
        parsed = parse_sales_register(content, name, ret.period, gstin.gstin, rules, _gst_rate_reference_rules(tenant),
            known_profiles=_known_sales_format_profiles(tenant, gstin.id, name))
    except Exception as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    fatal = [issue for issue in parsed["issues"] if issue["severity"] == "ERROR"]
    if fatal or not parsed["rows"]:
        raise HTTPException(status_code=422, detail={"message": "Fix invalid rows in the books register before reconciling.", "issues": fatal})
    prepared = tenant.query(models.GSTR1SalesInvoice).filter_by(upload_id=ret.source_upload_id).all()
    def key(row):
        get = row.get if isinstance(row, dict) else lambda field: getattr(row, field)
        party = get("customer_gstin") or re.sub(r"[^a-z0-9]+", "", (get("customer_name") or "").casefold())
        return (party, normalize_invoice_number(get("invoice_number")))
    book_by_key, prep_by_key = {}, {}
    duplicates = []
    for row in parsed["rows"]:
        k = key(row)
        if k in book_by_key:
            duplicates.append({"invoice_number": row["invoice_number"], "customer_gstin": row["customer_gstin"], "reason": "Repeated in books register"})
        else:
            book_by_key[k] = row
    for row in prepared:
        prep_by_key.setdefault(key(row), row)
    cancelled_in_books = []
    for k, b in book_by_key.items():
        p = prep_by_key.get(k)
        if p and str(b.get("invoice_status", "ACTIVE")).upper() in {"CANCELLED", "CANCELED", "VOID"} and str(p.invoice_status).upper() not in {"CANCELLED", "CANCELED", "VOID"}:
            cancelled_in_books.append({"invoice_number": p.invoice_number, "message": "Books mark this invoice cancelled, but the GSTR-1 preparation contains it as active."})
    missing = [{"invoice_number": row["invoice_number"], "customer_name": row["customer_name"], "taxable_value": float(row["taxable_value"]), "invoice_value": float(row["invoice_value"])}
        for k, row in book_by_key.items() if k not in prep_by_key]
    extra = [{"invoice_number": row.invoice_number, "customer_name": row.customer_name, "taxable_value": float(row.taxable_value), "invoice_value": float(row.invoice_value)}
        for k, row in prep_by_key.items() if k not in book_by_key]
    mismatches = []
    for k in book_by_key.keys() & prep_by_key.keys():
        b, p = book_by_key[k], prep_by_key[k]
        b_tax = Decimal(str(b["taxable_value"])); p_tax = Decimal(str(p.taxable_value))
        b_total = Decimal(str(b["invoice_value"])); p_total = Decimal(str(p.invoice_value))
        b_gst = sum((Decimal(str(b[x])) for x in ("cgst", "sgst", "igst", "cess")), Decimal("0"))
        p_gst = p.cgst + p.sgst + p.igst + p.cess
        if (b_tax, b_gst, b_total) != (p_tax, p_gst, p_total):
            mismatches.append({"invoice_number": p.invoice_number, "books_taxable": float(b_tax), "preparation_taxable": float(p_tax),
                "books_tax": float(b_gst), "preparation_tax": float(p_gst), "books_total": float(b_total), "preparation_total": float(p_total)})
    book_totals = {"taxable_value": sum((Decimal(str(r["taxable_value"])) for r in parsed["rows"]), Decimal("0")),
        "tax": sum((sum((Decimal(str(r[k])) for k in ("cgst", "sgst", "igst", "cess")), Decimal("0")) for r in parsed["rows"]), Decimal("0")),
        "invoice_value": sum((Decimal(str(r["invoice_value"])) for r in parsed["rows"]), Decimal("0"))}
    prep_totals = {"taxable_value": sum((r.taxable_value for r in prepared), Decimal("0")),
        "tax": sum((r.cgst+r.sgst+r.igst+r.cess for r in prepared), Decimal("0")),
        "invoice_value": sum((r.invoice_value for r in prepared), Decimal("0"))}
    result = {"source": "Separately uploaded books register", "filename": name, "sha256": parsed["source_sha256"],
        "compared_at": datetime.utcnow().isoformat()+"Z", "books_invoice_count": len(parsed["rows"]), "preparation_invoice_count": len(prepared),
        "books_totals": {k: float(v) for k,v in book_totals.items()}, "preparation_totals": {k: float(v) for k,v in prep_totals.items()},
        "totals_difference_books_minus_preparation": {k: float(book_totals[k]-prep_totals[k]) for k in book_totals},
        "missing_from_preparation": missing, "not_in_books_register": extra, "amount_mismatches": mismatches,
        "duplicate_books_rows": duplicates, "cancelled_in_books_but_active_in_preparation": cancelled_in_books}
    ret.reconciliation_result = result
    audit.record(tenant.db, firm_id=tenant.firm_id, user_id=actor.id, action=AuditAction.GSTR1_BOOKS_RECONCILIATION_RECORDED,
        entity_type="GSTR1Return", entity_id=ret.id, metadata={"period": ret.period, "source": name,
        "missing": len(missing), "extra": len(extra), "mismatched": len(mismatches), "sha256": parsed["source_sha256"]},
        ip_address=client_ip(request), user_agent=user_agent(request))
    tenant.commit()
    return result


@router.get("/gstins/{gstin_id}/uploads")
def list_sales_upload_versions(gstin_id: str, period: Optional[str] = Query(None),
        actor: models.User = Depends(require_role(Role.CLIENT_VIEWER)), tenant: TenantSession = Depends(get_tenant)):
    gstin = tenant.get_or_404(models.ClientGSTIN, gstin_id, "GSTIN not found")
    assert_client_visible(actor, gstin.client_id)
    query = tenant.query(models.GSTR1SalesUpload).filter_by(client_gstin_id=gstin_id)
    if period:
        query = query.filter_by(period=_period(period))
    rows = query.order_by(models.GSTR1SalesUpload.period.desc(), models.GSTR1SalesUpload.version.desc()).all()
    return [{"id": row.id, "period": row.period, "version": row.version, "is_current": row.is_current,
        "original_filename": row.original_filename, "source_sha256": row.source_sha256,
        "row_count": row.row_count, "issue_count": row.issue_count, "uploaded_by": row.uploaded_by,
        "uploaded_at": row.uploaded_at.isoformat()} for row in rows]


@router.get("/uploads/{upload_id}/source")
def download_sales_upload_source(upload_id: str, request: Request,
        actor: models.User = Depends(require_role(Role.ASSOCIATE)), tenant: TenantSession = Depends(get_tenant)):
    upload = tenant.get_or_404(models.GSTR1SalesUpload, upload_id, "Sales register upload not found")
    if not upload.file_path:
        raise HTTPException(status_code=404, detail="The source file is not available.")
    try:
        content = get_storage().load(upload.file_path, firm_id=tenant.firm_id)
    except Exception as exc:
        raise HTTPException(status_code=404, detail="The encrypted source file could not be retrieved.") from exc
    filename = (upload.original_filename or "sales-register.csv").replace("\r", "").replace("\n", "").replace('"', "")
    audit.record(tenant.db, firm_id=tenant.firm_id, user_id=actor.id, action=AuditAction.GSTR1_SOURCE_REGISTER_DOWNLOADED,
        entity_type="GSTR1SalesUpload", entity_id=upload.id, metadata={"period": upload.period, "version": upload.version},
        ip_address=client_ip(request), user_agent=user_agent(request))
    tenant.commit()
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" if filename.lower().endswith(".xlsx") else "text/csv"
    return Response(content=content, media_type=media_type,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename, safe='')}"})


@router.get("/returns/{return_id}/checklist")
def get_return_checklist(return_id: str, actor: models.User = Depends(require_role(Role.CLIENT_VIEWER)),
                         tenant: TenantSession = Depends(get_tenant)):
    ret = tenant.get_or_404(models.GSTR1Return, return_id, "GSTR-1 return not found")
    gstin = tenant.get_or_404(models.ClientGSTIN, ret.client_gstin_id, "GSTIN not found")
    assert_client_visible(actor, gstin.client_id)
    result = _get_return_data(gstin.id, ret.period, actor, tenant)
    invoices = tenant.query(models.GSTR1SalesInvoice).filter_by(upload_id=ret.source_upload_id).all()
    issue_rows = tenant.query(models.GSTR1ValidationIssue).filter_by(upload_id=ret.source_upload_id).all()
    documents = result["documents_issued"]
    needs_table13 = ret.period >= "2025-05" and bool(invoices)
    source_taxable = sum((row.taxable_value * (-1 if normalize_document_type(row.document_type).startswith("CREDIT_NOTE") else 1)
        for row in invoices), start=Decimal("0.00"))
    source_tax = sum(((row.cgst + row.sgst + row.igst + row.cess) *
        (-1 if normalize_document_type(row.document_type).startswith("CREDIT_NOTE") else 1)
        for row in invoices), start=Decimal("0.00"))
    hsn_taxable = sum((Decimal(str(row["taxable_value"])) for row in result["hsn_summary"].values()), start=Decimal("0.00"))
    hsn_tax = sum((Decimal(str(row["cgst"])) + Decimal(str(row["sgst"])) + Decimal(str(row["igst"])) + Decimal(str(row["cess"]))
                   for row in result["hsn_summary"].values()), start=Decimal("0.00"))
    grouped_totals_match = (len(invoices) == sum(row["record_count"] for row in result["hsn_summary"].values())
        and source_taxable.quantize(Decimal("0.01")) == hsn_taxable.quantize(Decimal("0.01"))
        and source_tax.quantize(Decimal("0.01")) == hsn_tax.quantize(Decimal("0.01")))
    note_invoice_ids = {row.id for row in invoices if normalize_document_type(row.document_type).startswith(("CREDIT_NOTE", "DEBIT_NOTE"))}
    note_review_issues = [issue for issue in issue_rows if issue.invoice_id in note_invoice_ids and issue.code.startswith("NOTE_")]
    open_note_checks = any(issue.status == "OPEN" for issue in note_review_issues)
    portal_confirmed = {issue.invoice_id for issue in note_review_issues
        if issue.code == "NOTE_PORTAL_REPORTING_CHECK" and issue.status == "RESOLVED" and issue.review_decision == "ACCEPT_VALID"}
    note_reporting_confirmed = note_invoice_ids.issubset(portal_confirmed)
    note_check_status = ("NOT_APPLICABLE" if not note_invoice_ids else
        "BLOCKED" if open_note_checks or not note_reporting_confirmed else "PASS")
    note_check_detail = ("No credit or debit notes were imported." if not note_invoice_ids else
        "Resolve original-invoice, amount, duplicate and tax checks, then confirm each note was checked against filed returns. Portal history is not connected, so the accountant must verify it using filing records.")
    unresolved_categories = sum(1 for row in result["invoices"] if row["category"] in {"UNCLASSIFIED", "OTHER"})
    checks = [
        {"key": "source_register", "label": "Sales register imported", "status": "PASS" if bool(invoices) else "BLOCKED", "detail": f"{len(invoices)} valid invoice rows imported."},
        {"key": "source_rows", "label": "Source rows accepted", "status": "BLOCKED" if result["upload_issues"] else "PASS", "detail": f"{len(result['upload_issues'])} invalid source rows remain excluded."},
        {"key": "classification", "label": "Every supply has a reporting classification", "status": "BLOCKED" if unresolved_categories else "PASS", "detail": f"{unresolved_categories} records need their document type or recipient classification reviewed before filing."},
        {"key": "manager_review", "label": "Invoice checks reviewed by a manager", "status": "PASS" if invoices and all(row.status == models.GSTR1InvoiceStatus.REVIEWED for row in invoices) else "BLOCKED", "detail": "Every imported invoice requires an independent manager review."},
        {"key": "exceptions", "label": "Validation issues resolved", "status": "PASS" if result["summary"]["open_issues"] == 0 else "BLOCKED", "detail": f"{result['summary']['open_issues']} open review items."},
        {"key": "hsn_rollup", "label": "B2B/B2C HSN summary ties to imported lines", "status": "PASS" if grouped_totals_match else "BLOCKED", "detail": f"Source taxable ₹{source_taxable:,.2f}; HSN grouped taxable ₹{hsn_taxable:,.2f}."},
        {"key": "hsn_rules", "label": "HSN/SAC checked against effective-dated GST reference", "status": "BLOCKED", "detail": "An authoritative, period-effective HSN/SAC validation catalogue is not configured."},
        {"key": "gst_rate_rules", "label": "GST rate checked against effective-dated GST reference", "status": "BLOCKED" if not result["gst_rate_reference"]["applicable"] else "PASS", "detail": result["gst_rate_reference"]["detail"]},
        {"key": "table13", "label": "Table 13 documents-issued records reviewed", "status": ("PASS" if not needs_table13 or (documents and all(row["reviewed_by"] for row in documents)) else "BLOCKED"), "detail": "A current Table 13 series count must be reviewed by a different user."},
        {"key": "note_linkage", "label": "Credit/debit note relationships checked", "status": note_check_status, "detail": note_check_detail},
        {"key": "ecommerce", "label": "E-commerce supplies applicability reviewed", "status": "BLOCKED", "detail": "Tables 14/15 applicability has not been captured in this version."},
        {"key": "external_reconciliation", "label": "Books sales register reconciled to this preparation", "status": ("NOT_APPLICABLE" if not invoices else "BLOCKED" if not ret.reconciliation_result else "PASS" if not (ret.reconciliation_result.get("missing_from_preparation") or ret.reconciliation_result.get("not_in_books_register") or ret.reconciliation_result.get("amount_mismatches") or ret.reconciliation_result.get("duplicate_books_rows") or ret.reconciliation_result.get("cancelled_in_books_but_active_in_preparation")) else "BLOCKED"), "detail": "Upload a separate books sales register to compare invoice coverage and totals." if invoices and not ret.reconciliation_result else ("Reconciliation differences still need review." if ret.reconciliation_result and (ret.reconciliation_result.get("missing_from_preparation") or ret.reconciliation_result.get("not_in_books_register") or ret.reconciliation_result.get("amount_mismatches") or ret.reconciliation_result.get("duplicate_books_rows") or ret.reconciliation_result.get("cancelled_in_books_but_active_in_preparation")) else "The supplied books register ties to the current GSTR-1 preparation." if ret.reconciliation_result else "No invoices are imported." )},
        {"key": "partner_approval", "label": "Partner approval recorded", "status": "PASS" if ret.approved_by else "BLOCKED", "detail": "Partner approval covers this preparation package only; filing remains a separate human action."},
    ]
    ready = all(item["status"] in {"PASS", "NOT_APPLICABLE"} for item in checks)
    return {"return_id": ret.id, "period": ret.period, "return_status": ret.status.value,
        "readiness": "READY_FOR_FILING" if ready else "REVIEW_REQUIRED", "ready_for_filing": ready,
        "filed": False, "checks": checks}


class InvoiceReviewIn(BaseModel):
    decision: str
    note: Optional[str] = None


class IssueDecisionIn(BaseModel):
    decision: str
    note: Optional[str] = None


@router.post("/issues/{issue_id}/decision")
def decide_validation_issue(issue_id: str, payload: IssueDecisionIn, request: Request,
        actor: models.User = Depends(require_exact_roles(Role.MANAGER, Role.OWNER)), tenant: TenantSession = Depends(get_tenant)):
    """Record an accountable decision. Corrections stay open until corrected source data is re-imported."""
    allowed = {"ACCEPT_VALID", "IGNORE", "NEEDS_CORRECTION", "REOPEN"}
    if payload.decision not in allowed:
        raise HTTPException(status_code=422, detail=f"Decision must be one of: {', '.join(sorted(allowed))}.")
    issue = tenant.get_or_404(models.GSTR1ValidationIssue, issue_id, "Validation item not found")
    upload = tenant.get_or_404(models.GSTR1SalesUpload, issue.upload_id, "Sales upload not found")
    if upload.uploaded_by == actor.id:
        raise HTTPException(status_code=409, detail="The preparer cannot review their own imported invoices.")
    ret = tenant.query(models.GSTR1Return).filter_by(client_gstin_id=upload.client_gstin_id, period=upload.period).first()
    if not ret or ret.status != models.GSTR1ReturnStatus.DRAFT or ret.source_upload_id != upload.id:
        raise HTTPException(status_code=409, detail="Only checks in the active draft can be reviewed.")
    issue.review_decision = payload.decision
    issue.review_note = (payload.note or "").strip()[:1000] or None
    issue.reviewed_by = actor.id
    issue.reviewed_at = datetime.utcnow()
    issue.status = "RESOLVED" if payload.decision in {"ACCEPT_VALID", "IGNORE"} else "OPEN"
    if issue.invoice_id:
        invoice = tenant.get(models.GSTR1SalesInvoice, issue.invoice_id)
        other_open_for_invoice = (tenant.query(models.GSTR1ValidationIssue).filter_by(invoice_id=issue.invoice_id, status="OPEN")
            .filter(models.GSTR1ValidationIssue.id != issue.id).count())
        if invoice and not (other_open_for_invoice or issue.status == "OPEN"):
            invoice.status = models.GSTR1InvoiceStatus.REVIEWED
            invoice.reviewed_by = actor.id
            invoice.reviewed_at = datetime.utcnow()
        elif invoice:
            invoice.status = models.GSTR1InvoiceStatus.NEEDS_REVIEW
            invoice.reviewed_by = None
            invoice.reviewed_at = None
    audit.record(tenant.db, firm_id=tenant.firm_id, user_id=actor.id, action=AuditAction.GSTR1_INVOICE_REVIEWED,
        entity_type="GSTR1ValidationIssue", entity_id=issue.id,
        metadata={"decision": payload.decision, "note": issue.review_note, "code": issue.code},
        ip_address=client_ip(request), user_agent=user_agent(request))
    tenant.commit()
    return {"id": issue.id, "status": issue.status, "decision": issue.review_decision, "review_note": issue.review_note}


class DocumentIssuedIn(BaseModel):
    document_type: str
    series_from: str
    series_to: str
    issued_count: int
    cancelled_count: int = 0


@router.get("/returns/{return_id}/documents-issued")
def list_documents_issued(return_id: str, actor: models.User = Depends(require_role(Role.CLIENT_VIEWER)),
                          tenant: TenantSession = Depends(get_tenant)):
    ret = tenant.get_or_404(models.GSTR1Return, return_id, "GSTR-1 return not found")
    gstin = tenant.get_or_404(models.ClientGSTIN, ret.client_gstin_id, "GSTIN not found")
    assert_client_visible(actor, gstin.client_id)
    rows = tenant.query(models.GSTR1DocumentIssued).filter_by(return_id=ret.id, is_current=True).all()
    return [{"id": r.id, "document_type": r.document_type, "series_from": r.series_from,
        "series_to": r.series_to, "issued_count": r.issued_count, "cancelled_count": r.cancelled_count,
        "net_issued": r.issued_count-r.cancelled_count, "version": r.version,
        "prepared_by": r.prepared_by, "reviewed_by": r.reviewed_by,
        "reviewed_at": r.reviewed_at.isoformat() if r.reviewed_at else None} for r in rows]


@router.post("/returns/{return_id}/documents-issued")
def record_documents_issued(return_id: str, payload: DocumentIssuedIn, request: Request,
        actor: models.User = Depends(require_role(Role.ASSOCIATE)), tenant: TenantSession = Depends(get_tenant)):
    ret = tenant.get_or_404(models.GSTR1Return, return_id, "GSTR-1 return not found")
    if ret.status != models.GSTR1ReturnStatus.DRAFT:
        raise HTTPException(status_code=409, detail="Document counts can only be changed while the return is a draft.")
    doc_type = payload.document_type.strip().upper()
    if doc_type not in {"INVOICE", "CREDIT_NOTE", "DEBIT_NOTE", "OTHER"}:
        raise HTTPException(status_code=422, detail="Document type must be Invoice, Credit note, Debit note, or Other.")
    if not doc_type or len(doc_type) > 40 or not payload.series_from.strip() or not payload.series_to.strip():
        raise HTTPException(status_code=422, detail="Document type and series range are required.")
    if payload.issued_count < 0 or payload.cancelled_count < 0 or payload.cancelled_count > payload.issued_count:
        raise HTTPException(status_code=422, detail="Counts must be non-negative, and cancelled documents cannot exceed issued documents.")
    existing = tenant.query(models.GSTR1DocumentIssued).filter_by(return_id=ret.id, document_type=doc_type,
        series_from=payload.series_from.strip(), is_current=True).first()
    version = 1
    if existing:
        version = existing.version + 1
        existing.is_current = False
    record = models.GSTR1DocumentIssued(return_id=ret.id, document_type=doc_type,
        series_from=payload.series_from.strip(), series_to=payload.series_to.strip(),
        issued_count=payload.issued_count, cancelled_count=payload.cancelled_count,
        version=version, is_current=True, prepared_by=actor.id)
    tenant.add(record)
    audit.record(tenant.db, firm_id=tenant.firm_id, user_id=actor.id, action=AuditAction.GSTR1_DOCUMENTS_ISSUED_RECORDED,
        entity_type="GSTR1Return", entity_id=ret.id, metadata={"document_type": doc_type,
        "series_from": payload.series_from, "series_to": payload.series_to,
        "issued_count": payload.issued_count, "cancelled_count": payload.cancelled_count, "version": version},
        ip_address=client_ip(request), user_agent=user_agent(request))
    tenant.commit()
    return {"id": record.id, "document_type": record.document_type, "version": record.version, "status": "DRAFT"}


@router.post("/documents-issued/{record_id}/review")
def review_documents_issued(record_id: str, request: Request,
        actor: models.User = Depends(require_exact_roles(Role.MANAGER, Role.OWNER)), tenant: TenantSession = Depends(get_tenant)):
    record = tenant.get_or_404(models.GSTR1DocumentIssued, record_id, "Document record not found")
    if not record.is_current:
        raise HTTPException(status_code=409, detail="Only the current document count version can be reviewed.")
    if record.prepared_by == actor.id:
        raise HTTPException(status_code=409, detail="The preparer cannot review their own document count.")
    record.reviewed_by = actor.id
    record.reviewed_at = datetime.utcnow()
    audit.record(tenant.db, firm_id=tenant.firm_id, user_id=actor.id, action=AuditAction.GSTR1_DOCUMENTS_ISSUED_REVIEWED,
        entity_type="GSTR1DocumentIssued", entity_id=record.id, metadata={"version": record.version},
        ip_address=client_ip(request), user_agent=user_agent(request))
    tenant.commit()
    return {"id": record.id, "reviewed_by": record.reviewed_by, "reviewed_at": record.reviewed_at}


@router.post("/invoices/{invoice_id}/review")
def review_invoice(invoice_id: str, payload: InvoiceReviewIn, request: Request,
                   actor: models.User = Depends(require_exact_roles(Role.MANAGER, Role.OWNER)), tenant: TenantSession = Depends(get_tenant)):
    if payload.decision not in {"reviewed", "needs_review"}:
        raise HTTPException(status_code=422, detail="Decision must be reviewed or needs_review.")
    invoice = tenant.get_or_404(models.GSTR1SalesInvoice, invoice_id, "Invoice not found")
    upload = tenant.get_or_404(models.GSTR1SalesUpload, invoice.upload_id, "Sales upload not found")
    if upload.uploaded_by == actor.id:
        raise HTTPException(status_code=409, detail="The preparer cannot review their own imported invoices.")
    ret = tenant.query(models.GSTR1Return).filter_by(client_gstin_id=upload.client_gstin_id, period=upload.period).first()
    if not ret or ret.status != models.GSTR1ReturnStatus.DRAFT or ret.source_upload_id != upload.id:
        raise HTTPException(status_code=409, detail="Only invoices in the active draft can be reviewed.")
    if payload.decision == "reviewed" and tenant.query(models.GSTR1ValidationIssue).filter_by(invoice_id=invoice.id, status="OPEN").count():
        raise HTTPException(status_code=409, detail="Resolve each open check separately as accepted, ignored, or needing correction.")
    invoice.status = models.GSTR1InvoiceStatus.REVIEWED if payload.decision == "reviewed" else models.GSTR1InvoiceStatus.NEEDS_REVIEW
    invoice.reviewed_by = actor.id if payload.decision == "reviewed" else None
    invoice.reviewed_at = datetime.utcnow() if payload.decision == "reviewed" else None
    issues = tenant.query(models.GSTR1ValidationIssue).filter_by(invoice_id=invoice.id, status="OPEN").all()
    for issue in issues:
        issue.status = "REVIEWED" if payload.decision == "reviewed" else "OPEN"
        issue.reviewed_by = actor.id if payload.decision == "reviewed" else None
        issue.reviewed_at = datetime.utcnow() if payload.decision == "reviewed" else None
    audit.record(tenant.db, firm_id=tenant.firm_id, user_id=actor.id, action=AuditAction.GSTR1_INVOICE_REVIEWED,
        entity_type="GSTR1SalesInvoice", entity_id=invoice.id,
        metadata={"decision": payload.decision, "note": (payload.note or "")[:1000]},
        ip_address=client_ip(request), user_agent=user_agent(request))
    tenant.commit()
    return {"id": invoice.id, "status": invoice.status.value, "reviewed_by": invoice.reviewed_by, "reviewed_at": invoice.reviewed_at}


@router.post("/returns/{return_id}/submit-review")
def submit_return_for_partner_review(return_id: str, request: Request,
        actor: models.User = Depends(require_exact_roles(Role.MANAGER, Role.OWNER)), tenant: TenantSession = Depends(get_tenant)):
    ret = tenant.get_or_404(models.GSTR1Return, return_id, "GSTR-1 return not found")
    if ret.status != models.GSTR1ReturnStatus.DRAFT:
        raise HTTPException(status_code=409, detail="Only a draft return can be submitted for partner review.")
    if ret.created_by == actor.id:
        raise HTTPException(status_code=409, detail="The return preparer cannot submit their own work for approval.")
    upload = tenant.get_or_404(models.GSTR1SalesUpload, ret.source_upload_id, "Sales register not found")
    if upload.uploaded_by == actor.id:
        raise HTTPException(status_code=409, detail="The register uploader cannot review this return.")
    invoices = tenant.query(models.GSTR1SalesInvoice).filter_by(upload_id=upload.id).all()
    if not invoices or any(row.status != models.GSTR1InvoiceStatus.REVIEWED for row in invoices):
        raise HTTPException(status_code=409, detail="Every imported invoice must be reviewed before partner review.")
    open_issues = tenant.query(models.GSTR1ValidationIssue).filter_by(upload_id=upload.id, status="OPEN").count()
    if open_issues:
        raise HTTPException(status_code=409, detail="Open validation items must be reviewed before partner review.")
    if ret.period >= "2025-05" and invoices:
        documents = tenant.query(models.GSTR1DocumentIssued).filter_by(return_id=ret.id, is_current=True).all()
        if not documents or any(not row.reviewed_by for row in documents):
            raise HTTPException(status_code=409, detail="Table 13 document counts must be prepared and independently reviewed for this period.")
    ret.status = models.GSTR1ReturnStatus.IN_REVIEW
    ret.reviewed_by = actor.id
    ret.reviewed_at = datetime.utcnow()
    audit.record(tenant.db, firm_id=tenant.firm_id, user_id=actor.id, action=AuditAction.GSTR1_SUBMITTED_FOR_PARTNER_REVIEW,
        entity_type="GSTR1Return", entity_id=ret.id, metadata={"period": ret.period},
        ip_address=client_ip(request), user_agent=user_agent(request))
    tenant.commit()
    return {"id": ret.id, "status": ret.status.value, "reviewed_by": actor.id}


@router.post("/returns/{return_id}/approve")
def approve_return(return_id: str, request: Request,
        actor: models.User = Depends(require_exact_roles(Role.PARTNER, Role.OWNER)), tenant: TenantSession = Depends(get_tenant)):
    ret = tenant.get_or_404(models.GSTR1Return, return_id, "GSTR-1 return not found")
    if ret.status != models.GSTR1ReturnStatus.IN_REVIEW:
        raise HTTPException(status_code=409, detail="Only a manager-reviewed return can be approved.")
    if actor.id in {ret.created_by, ret.reviewed_by}:
        raise HTTPException(status_code=409, detail="The preparer or reviewer cannot approve this return.")
    ret.status = models.GSTR1ReturnStatus.APPROVED
    ret.approved_by = actor.id
    ret.approved_at = datetime.utcnow()
    audit.record(tenant.db, firm_id=tenant.firm_id, user_id=actor.id, action=AuditAction.GSTR1_RETURN_PARTNER_APPROVED,
        entity_type="GSTR1Return", entity_id=ret.id, metadata={"period": ret.period},
        ip_address=client_ip(request), user_agent=user_agent(request))
    tenant.commit()
    return {"id": ret.id, "status": ret.status.value, "approved_by": actor.id, "approved_at": ret.approved_at}


@router.get("/returns/{return_id}/approved-export")
def approved_export(return_id: str, request: Request,
                    actor: models.User = Depends(require_exact_roles(Role.PARTNER, Role.OWNER)),
                    tenant: TenantSession = Depends(get_tenant)):
    ret = tenant.get_or_404(models.GSTR1Return, return_id, "GSTR-1 return not found")
    if ret.status != models.GSTR1ReturnStatus.APPROVED or not ret.approved_by:
        raise HTTPException(status_code=409, detail="Partner approval is required before generating an approved export.")
    gstin = tenant.get_or_404(models.ClientGSTIN, ret.client_gstin_id, "GSTIN not found")
    assert_client_visible(actor, gstin.client_id)
    result = _get_return_data(gstin.id, ret.period, actor, tenant)
    result["status"] = models.GSTR1ReturnStatus.APPROVED.value
    result["approved_by"] = ret.approved_by
    result["approved_at"] = ret.approved_at.isoformat() if ret.approved_at else None
    audit.record(tenant.db, firm_id=tenant.firm_id, user_id=actor.id, action=AuditAction.GSTR1_APPROVED_EXPORT_GENERATED,
        entity_type="GSTR1Return", entity_id=ret.id, metadata={"period": ret.period, "approved_by": ret.approved_by},
        ip_address=client_ip(request), user_agent=user_agent(request))
    tenant.commit()
    return {"document_status": "PARTNER APPROVED PREPARATION DATA — NOT FILED", "return": result}


@router.get("/gstins/{gstin_id}/returns/{period}/draft.xlsx")
def export_draft_workbook(gstin_id: str, period: str, request: Request,
        actor: models.User = Depends(require_role(Role.ASSOCIATE)), tenant: TenantSession = Depends(get_tenant)):
    """Internal accountant workbook; it is not a GST Portal upload workbook."""
    result = _get_return_data(gstin_id, period, actor, tenant)
    if result["status"] == models.GSTR1ReturnStatus.APPROVED.value:
        raise HTTPException(status_code=409, detail="This return is partner approved; use its approved preparation export.")
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame([{"Document status": "DRAFT - REVIEW REQUIRED - NOT FILED",
                       "Period": period, "Rule version": result["rule_version"],
                       "Source upload version": result["upload"]["version"] if result["upload"] else None,
                       "Source file": result["upload"]["filename"] if result["upload"] else None}]).to_excel(writer, sheet_name="Read Me", index=False)
        pd.DataFrame(result["invoices"]).drop(columns=["issues"], errors="ignore").to_excel(writer, sheet_name="Invoice Data", index=False)
        pd.DataFrame([{"HSN/SAC": item["hsn_sac"], "Reporting group": item["reporting_group"], **item}
                      for item in result["hsn_summary"].values()]).to_excel(writer, sheet_name="HSN Summary", index=False)
        pd.DataFrame(result["documents_issued"]).to_excel(writer, sheet_name="Documents Issued", index=False)
        pd.DataFrame(result["upload_issues"]).to_excel(writer, sheet_name="Unimported Rows", index=False)
    output.seek(0)
    audit.record(tenant.db, firm_id=tenant.firm_id, user_id=actor.id, action=AuditAction.GSTR1_DRAFT_EXPORTED,
        entity_type="GSTR1Return", entity_id=result["id"], metadata={"period": period, "format": "XLSX", "document_status": "DRAFT"},
        ip_address=client_ip(request), user_agent=user_agent(request))
    tenant.commit()
    headers = {"Content-Disposition": f'attachment; filename="GSTR1-{period}-draft.xlsx"'}
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)


@router.get("/gstins/{gstin_id}/returns/{period}/annexure.xlsx")
def export_annexure_workbook(gstin_id: str, period: str, request: Request,
        actor: models.User = Depends(require_role(Role.ASSOCIATE)), tenant: TenantSession = Depends(get_tenant)):
    """Download a turnover annexure; supplementary export columns are breakouts, not additions."""
    result = _get_return_data(gstin_id, period, actor, tenant)
    report = result["annexure_report"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Annexure"
    headers = ["Particulars", "Records", "Taxable Amount", "CGST", "SGST", "IGST", "Cess", "Export/SEZ", "IGST-Export"]
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    sheet.cell(1, 1, report["period_label"])
    sheet.cell(1, 1).font = Font(bold=True, size=14, color="17365D")
    sheet.cell(1, 1).fill = PatternFill("solid", fgColor="DDEBF7")
    sheet.cell(1, 1).alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 26
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(2, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="left" if col == 1 else "right")
    amount_fields = ["taxable_amount", "cgst", "sgst", "igst", "cess", "export_sez", "igst_export"]
    for row_num, row in enumerate(report["rows"], start=3):
        values = [row["particulars"], row["records"], *[float(row[field]) for field in amount_fields]]
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row_num, col, value)
            if col >= 3:
                cell.number_format = '₹#,##0.00;[Red]-₹#,##0.00;-'
                cell.alignment = Alignment(horizontal="right")
            if row["records"] and row["particulars"] == "Needs category review":
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
    total_row = len(report["rows"]) + 3
    total_values = ["Net turnover", report["totals"]["records"], *[float(report["totals"][field]) for field in amount_fields]]
    for col, value in enumerate(total_values, start=1):
        cell = sheet.cell(total_row, col, value)
        cell.font = Font(bold=True, color="17365D")
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
        if col >= 3:
            cell.number_format = '₹#,##0.00;[Red]-₹#,##0.00;-'
    note_row = total_row + 2
    sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=len(headers))
    sheet.cell(note_row, 1, report["note"])
    sheet.cell(note_row, 1).font = Font(italic=True, color="666666", size=10)
    sheet.cell(note_row, 1).alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[note_row].height = 30
    for col, width in enumerate([38, 12, 20, 18, 18, 18, 16, 20, 20], start=1):
        sheet.column_dimensions[chr(64 + col)].width = width
    sheet.freeze_panes = "C3"
    sheet.auto_filter.ref = f"A2:I{total_row - 1}"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    audit.record(tenant.db, firm_id=tenant.firm_id, user_id=actor.id, action=AuditAction.GSTR1_DRAFT_EXPORTED,
        entity_type="GSTR1Return", entity_id=result["id"], metadata={"period": period, "format": "ANNEXURE_XLSX"},
        ip_address=client_ip(request), user_agent=user_agent(request))
    tenant.commit()
    headers = {"Content-Disposition": f'attachment; filename="GSTR1-{period}-annexure.xlsx"'}
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)


def _annexure_export_rows(gstin_id: str, period: str, actor: models.User, tenant: TenantSession):
    result = _get_return_data(gstin_id, period, actor, tenant)
    gstin = tenant.get_or_404(models.ClientGSTIN, gstin_id, "GSTIN not found")
    upload = tenant.get(models.GSTR1SalesUpload, result["upload"]["id"]) if result["upload"] else None
    records = tenant.query(models.GSTR1SalesInvoice).filter_by(upload_id=upload.id).order_by(
        models.GSTR1SalesInvoice.invoice_date, models.GSTR1SalesInvoice.invoice_number).all() if upload else []
    record_ids = [row.id for row in records]
    issue_rows = tenant.query(models.GSTR1ValidationIssue).filter(models.GSTR1ValidationIssue.invoice_id.in_(record_ids)).all() if record_ids else []
    by_record = defaultdict(list)
    for issue in issue_rows:
        by_record[issue.invoice_id].append(issue)
    rules = [{"version": rule.version, "definition": rule.definition, "effective_from": rule.effective_from.isoformat(),
        "effective_to": rule.effective_to.isoformat() if rule.effective_to else None} for rule in tenant.db.query(models.GSTR1RuleDefinition)
        .filter(models.GSTR1RuleDefinition.rule_key == "B2CL_THRESHOLD").all()]
    relationship_by_id = {item["id"]: item.get("noteRelationship") for item in result["invoices"]}
    rows = [transaction_view(row, by_record[row.id], gstin.gstin, period, rules) for row in records]
    for row, record in zip(rows, records):
        relationship = relationship_by_id.get(record.id)
        original = (relationship or {}).get("original") or {}
        if relationship:
            row.update({"noteLinkStatus": relationship.get("status"),
                "noteRelationship": relationship,
                "originalInvoiceNumber": original.get("invoice_number") or row.get("originalInvoiceNumber"),
                "originalInvoiceDate": original.get("invoice_date") or row.get("originalInvoiceDate"),
                "originalInvoiceSource": original.get("source"),
                "originalInvoiceValue": original.get("invoice_value"),
                "originalInvoiceTaxableValue": original.get("taxable_value"),
                "originalInvoiceTax": original.get("tax"),
                "cumulativeCreditNotes": original.get("cumulative_credit_notes"),
                "remainingOriginalTaxableValue": original.get("remaining_taxable_value"),
                "portalFilingCheck": "Manual check required; GST Portal filing history is not connected"})
    return result, records, rows


@router.get("/gstins/{gstin_id}/returns/{period}/annexure-workbook.xlsx")
def export_full_annexure_workbook(gstin_id: str, period: str, request: Request,
        actor: models.User = Depends(require_role(Role.ASSOCIATE)), tenant: TenantSession = Depends(get_tenant)):
    """Multi-sheet working paper from canonical, normalized return records."""
    period = _period(period)
    result, records, rows = _annexure_export_rows(gstin_id, period, actor, tenant)
    labels = {"invoiceNumber":"Invoice Number", "invoiceDate":"Invoice Date", "customerName":"Customer Name",
        "customerGstin":"Customer GSTIN", "placeOfSupply":"Place of Supply", "supplyType":"Supply Type",
        "documentType":"Document Type", "taxableValue":"Taxable Value", "gstRate":"GST Rate (%)", "cgst":"CGST",
        "sgst":"SGST", "igst":"IGST", "cess":"Cess", "invoiceValue":"Invoice Value", "hsnSac":"HSN/SAC",
        "classification":"GSTR-1 Classification", "gstr1Table":"GSTR-1 Table", "validationStatus":"Validation Status",
        "classificationReason":"Classification Reason", "quantity":"Quantity", "uqc":"UQC", "exportType":"Export Type",
        "shippingBillNumber":"Shipping Bill Number", "shippingBillDate":"Shipping Bill Date", "portCode":"Port Code",
        "originalInvoiceNumber":"Original Invoice Number", "originalInvoiceDate":"Original Invoice Date", "noteReason":"Note Reason",
        "noteLinkStatus":"Original Invoice Link Status", "originalInvoiceValue":"Original Invoice Value",
        "originalInvoiceTaxableValue":"Original Taxable Value", "originalInvoiceTax":"Original Tax",
        "cumulativeCreditNotes":"Cumulative Credit Notes", "remainingOriginalTaxableValue":"Remaining Original Taxable Value",
        "portalFilingCheck":"Filed Return Check"}
    def frame(selected):
        return pd.DataFrame([{labels[key]: value for key, value in row.items() if key in labels} for row in selected])
    hsn_rows = build_hsn_summary([{**row, "reportingGroup": row["classification"]} for row in rows])
    excluded = tenant.query(models.GSTR1ValidationIssue).filter_by(upload_id=result["upload"]["id"], invoice_id=None).all() if result["upload"] else []
    excluded_count = len({issue.row_number for issue in excluded})
    tax_sum = lambda field: sum((abs(float(getattr(row, field) or 0)) * (-1 if normalize_document_type(row.document_type).startswith("CREDIT_NOTE") else 1) for row in records), 0)
    all_exceptions = [{"Invoice Number": row["invoiceNumber"], "Customer": row["customerName"], "Issue": issue["title"],
        "Details": issue["message"], "Severity": issue["severity"], "Status": issue["status"]}
        for row in rows for issue in row["exceptions"]]
    all_exceptions.extend({"Invoice Number":"", "Customer":"", "Issue":issue.code.replace("_", " ").title(),
        "Details":issue.message, "Severity":issue.severity, "Status":issue.status, "Source Row":issue.row_number}
        for issue in excluded)
    related_ids = {result["id"]}
    if result["upload"]:
        related_ids.add(result["upload"]["id"])
    related_ids.update(row.id for row in records)
    related_logs = (tenant.query(models.AuditLog).filter(models.AuditLog.entity_id.in_(related_ids))
        .order_by(models.AuditLog.created_at.desc()).limit(500).all()) if related_ids else []
    actor_ids = {entry.user_id for entry in related_logs if entry.user_id}
    actor_names = {user.id: user.name for user in tenant.query(models.User).filter(models.User.id.in_(actor_ids)).all()} if actor_ids else {}
    audit_rows = [{"Date": entry.created_at.isoformat() + "Z", "Action": entry.action,
        "Performed by": actor_names.get(entry.user_id, "System"), "Record type": entry.entity_type or "—",
        "Details": json.dumps(entry.meta or {}, ensure_ascii=False, default=str)} for entry in related_logs]
    if not audit_rows:
        audit_rows = [{"Date": datetime.utcnow().isoformat() + "Z", "Action": "Annexure workbook generated",
            "Performed by": actor_names.get(actor.id, "Current user"), "Record type": "GSTR-1 return",
            "Details": json.dumps({"period": period})}]
    summary_frame = pd.DataFrame([{"Period": period, "Rule version": result["rule_version"],
        "GSTR-1 table mapping version": table_map_for_period(period)["version"],
        "Source rows": len(records) + excluded_count, "Processed": len(records), "Excluded": excluded_count,
        "Taxable Value": tax_sum("taxable_value"), "CGST": tax_sum("cgst"), "SGST": tax_sum("sgst"),
        "IGST": tax_sum("igst"), "Cess": tax_sum("cess")}])
    sheets = [("Summary", summary_frame), ("Sales Annexure", frame(rows)),
        ("B2B", frame([row for row in rows if row["classification"] == "B2B"])),
        ("B2C", frame([row for row in rows if row["classification"] in {"B2C Large", "B2C Others"}])),
        ("Exports", frame([row for row in rows if row["classification"] in {"Export", "SEZ supply", "Deemed export"}])),
        ("Credit_Debit_Notes", frame([row for row in rows if row["classification"] in {"Credit note", "Debit note"}])),
        ("HSN_Summary", pd.DataFrame([{"HSN/SAC": row["hsnSac"], "Description": row["description"], "Reporting category": row["reportingGroup"], "UQC": row["uqc"],
            "Quantity": row["quantity"], "GST Rate (%)": row["gstRate"], "Invoices": row["documents"], "Source records": row["records"],
            "Taxable Value": row["taxableValue"], "CGST": row["cgst"], "SGST": row["sgst"], "IGST": row["igst"], "Cess": row["cess"]} for row in hsn_rows])),
        ("Exceptions", pd.DataFrame(all_exceptions)),
        ("Reconciliation", pd.DataFrame([{"Source rows": len(records)+excluded_count, "Processed": len(records), "Excluded": excluded_count,
            "Annexure taxable value": tax_sum("taxable_value"), "GSTR-1 preparation taxable value": tax_sum("taxable_value"),
            "Difference": 0, "Basis": "Both totals use the same normalized draft records; this is not a GST Portal reconciliation."}])),
        ("Audit_Log", pd.DataFrame(audit_rows))]
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for name, sheet_frame in sheets:
            if name in {"B2B", "B2C", "Exports", "Credit_Debit_Notes", "Exceptions"} and sheet_frame.empty:
                continue
            sheet_frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            sheet.freeze_panes = "A2"
            if sheet_frame.shape[1]:
                sheet.auto_filter.ref = sheet.dimensions
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="1F4E78")
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                sheet.row_dimensions[1].height = 30
                for cells in sheet.columns:
                    width = min(48, max(12, max((len(str(cell.value or "")) for cell in cells[:100]), default=12) + 2))
                    sheet.column_dimensions[cells[0].column_letter].width = width
    output.seek(0)
    audit.record(tenant.db, firm_id=tenant.firm_id, user_id=actor.id, action=AuditAction.GSTR1_DRAFT_EXPORTED,
        entity_type="GSTR1Return", entity_id=result["id"], metadata={"period": period, "format": "ANNEXURE_WORKBOOK"},
        ip_address=client_ip(request), user_agent=user_agent(request))
    tenant.commit()
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="GSTR1-{period}-annexure-working-paper.xlsx"'})


@router.get("/gstins/{gstin_id}/returns/{period}/annexure-report.xlsx")
def export_annexure_report_xlsx(gstin_id: str, period: str, request: Request,
        report_type: str = Query(..., pattern="^(sales|b2b|b2c|exports|notes|hsn_b2b|hsn_summary)$"),
        actor: models.User = Depends(require_role(Role.ASSOCIATE)), tenant: TenantSession = Depends(get_tenant)):
    """Export one focused Annexure report as a separate Excel file."""
    period = _period(period)
    result, _records, rows = _annexure_export_rows(gstin_id, period, actor, tenant)
    categories = {
        "b2b": {"B2B"}, "b2c": {"B2C Large", "B2C Others"},
        "exports": {"Export", "SEZ supply", "Deemed export"},
        "notes": {"Credit note", "Debit note"},
    }
    if report_type in categories:
        selected = [row for row in rows if row["classification"] in categories[report_type]]
    elif report_type == "hsn_b2b":
        selected = [row for row in rows if row["classification"] == "B2B"]
    else:
        selected = rows

    if report_type in {"hsn_b2b", "hsn_summary"}:
        summary_rows = build_hsn_summary([{**row, "reportingGroup": row["classification"]} for row in selected])
        headers = ["HSN/SAC", "Description", "Reporting category", "UQC", "Quantity", "GST Rate (%)", "Invoices", "Source records",
            "Taxable Value", "CGST", "SGST", "IGST", "Cess", "Total Tax"]
        values = [[r["hsnSac"], r["description"], r["reportingGroup"], r["uqc"], r["quantity"], r["gstRate"], r["documents"], r["records"],
            r["taxableValue"], r["cgst"], r["sgst"], r["igst"], r["cess"], r["cgst"] + r["sgst"] + r["igst"] + r["cess"]] for r in summary_rows]
        if summary_rows:
            totals = {field: sum(row[field] for row in summary_rows) for field in ("taxableValue", "cgst", "sgst", "igst", "cess")}
            values.append(["Monetary totals", "", "", "", "", "", "", "",
                totals["taxableValue"], totals["cgst"], totals["sgst"], totals["igst"], totals["cess"],
                totals["cgst"] + totals["sgst"] + totals["igst"] + totals["cess"]])
        title = "HSN/SAC Summary - B2B" if report_type == "hsn_b2b" else "HSN/SAC Summary - All Categories"
    else:
        headers = ["Invoice Number", "Invoice Date", "Customer Name", "Customer GSTIN", "Document Type",
            "Classification", "Place of Supply", "Taxable Value", "GST Rate (%)", "CGST", "SGST", "IGST",
            "Cess", "Invoice Value", "HSN/SAC", "GSTR-1 Table", "Status"]
        if report_type == "notes":
            headers.extend(["Original Invoice Number", "Original Invoice Date", "Original Invoice Source",
                "Original Invoice Taxable Value", "Original Invoice Tax", "Cumulative Credit Notes",
                "Remaining Original Taxable Value", "Note Reason", "Original Invoice Link Status", "Filed Return Check", "Exceptions"])
        values = [[r["invoiceNumber"], r["invoiceDate"], r["customerName"], r["customerGstin"],
            r["documentType"], r["classification"], r["placeOfSupply"], r["taxableValue"], r["gstRate"],
            r["cgst"], r["sgst"], r["igst"], r["cess"], r["invoiceValue"], r["hsnSac"],
            r["gstr1Table"], r["validationStatus"]] + ([r.get("originalInvoiceNumber"), r.get("originalInvoiceDate"),
                r.get("originalInvoiceSource"),
                r.get("originalInvoiceTaxableValue"), r.get("originalInvoiceTax"), r.get("cumulativeCreditNotes"),
                r.get("remainingOriginalTaxableValue"), r.get("noteReason"), r.get("noteLinkStatus"),
                r.get("portalFilingCheck"), " | ".join(f"{issue['title']}: {issue['message']}" for issue in r["exceptions"]) ]
                if report_type == "notes" else []) for r in selected]
        title = {"sales":"Sales Annexure", "b2b":"B2B Annexure", "b2c":"B2C Annexure",
            "exports":"Export and SEZ Annexure", "notes":"Credit and Debit Note Annexure"}[report_type]

    frame = pd.DataFrame(values, columns=headers)
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name=report_type.replace("_", " ").title()[:31], index=False, startrow=2)
        sheet = writer.sheets[report_type.replace("_", " ").title()[:31]]
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(headers)))
        sheet.cell(1, 1, f"{title} - {period}")
        sheet.cell(1, 1).font = Font(bold=True, size=14, color="17365D")
        sheet.cell(1, 1).fill = PatternFill("solid", fgColor="DDEBF7")
        sheet.freeze_panes = "A4"
        if headers:
            sheet.auto_filter.ref = f"A3:{sheet.cell(3 + len(frame), len(headers)).coordinate}"
            for cell in sheet[3]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
            for column in sheet.columns:
                width = min(42, max(12, max((len(str(cell.value or "")) for cell in column[:100]), default=12) + 2))
                # The title rows contain merged cells. Resolve the column
                # letter from its numeric index without inspecting MergedCell.
                column_number = column[0].column
                column_letter = ""
                while column_number:
                    column_number, remainder = divmod(column_number - 1, 26)
                    column_letter = chr(65 + remainder) + column_letter
                sheet.column_dimensions[column_letter].width = width
    audit.record(tenant.db, firm_id=tenant.firm_id, user_id=actor.id, action=AuditAction.GSTR1_DRAFT_EXPORTED,
        entity_type="GSTR1Return", entity_id=result["id"],
        metadata={"period": period, "format": "ANNEXURE_REPORT_XLSX", "report_type": report_type},
        ip_address=client_ip(request), user_agent=user_agent(request))
    tenant.commit()
    output.seek(0)
    return StreamingResponse(iter([output.getvalue()]), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="GSTR1-{period}-{report_type}.xlsx"'})


@router.get("/gstins/{gstin_id}/returns/{period}/annexure.csv")
def export_annexure_csv(gstin_id: str, period: str, request: Request,
        actor: models.User = Depends(require_role(Role.ASSOCIATE)), tenant: TenantSession = Depends(get_tenant)):
    """Export the complete sales transaction annexure as a business-facing CSV."""
    period = _period(period)
    result, _records, rows = _annexure_export_rows(gstin_id, period, actor, tenant)
    labels = {"invoiceNumber":"Invoice Number", "invoiceDate":"Invoice Date", "customerName":"Customer Name",
        "customerGstin":"Customer GSTIN", "placeOfSupply":"Place of Supply", "supplyType":"Supply Type", "documentType":"Document Type",
        "taxableValue":"Taxable Value", "gstRate":"GST Rate (%)", "cgst":"CGST", "sgst":"SGST", "igst":"IGST", "cess":"Cess",
        "invoiceValue":"Invoice Value", "hsnSac":"HSN/SAC", "classification":"GSTR-1 Classification", "gstr1Table":"GSTR-1 Table",
        "validationStatus":"Validation Status", "classificationReason":"Classification Reason", "quantity":"Quantity", "uqc":"UQC",
        "exportType":"Export Type", "shippingBillNumber":"Shipping Bill Number", "shippingBillDate":"Shipping Bill Date",
        "portCode":"Port Code", "originalInvoiceNumber":"Original Invoice Number", "originalInvoiceDate":"Original Invoice Date",
        "noteReason":"Note Reason", "exceptions":"Exceptions"}
    stream = StringIO()
    writer = csv.DictWriter(stream, fieldnames=list(labels.values()), extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        output_row = {labels[key]: value for key, value in row.items() if key in labels and key != "exceptions"}
        output_row["Exceptions"] = " | ".join(f"{issue['title']}: {issue['message']}" for issue in row["exceptions"])
        writer.writerow(output_row)
    audit.record(tenant.db, firm_id=tenant.firm_id, user_id=actor.id, action=AuditAction.GSTR1_DRAFT_EXPORTED,
        entity_type="GSTR1Return", entity_id=result["id"], metadata={"period": period, "format": "ANNEXURE_CSV"},
        ip_address=client_ip(request), user_agent=user_agent(request))
    tenant.commit()
    stream.seek(0)
    return StreamingResponse(iter([stream.getvalue()]), media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="GSTR1-{period}-sales-annexure.csv"'})


@router.get("/gstins/{gstin_id}/returns/{period}/annexure.pdf")
def export_annexure_pdf(gstin_id: str, period: str, request: Request,
        actor: models.User = Depends(require_role(Role.ASSOCIATE)), tenant: TenantSession = Depends(get_tenant)):
    """Create a compact printable summary and exception overview."""
    period = _period(period)
    result, _records, rows = _annexure_export_rows(gstin_id, period, actor, tenant)
    totals = result["annexure_report"]["totals"]
    excluded = len({item.get("row_number") for item in result["upload_issues"] if item.get("row_number") is not None})
    lines = [f"GSTR-1 Annexure Summary - {period}", f"Preparation rule version: {result['rule_version']}", "",
        f"Source rows: {len(rows)+excluded}    Processed: {len(rows)}    Excluded: {excluded}",
        f"Transactions with open checks: {len({row['invoiceNumber'] for row in rows if row['validationStatus'] != 'READY'})}",
        f"Taxable value (net of credit notes): INR {float(totals['taxable_amount']):,.2f}",
        f"CGST: INR {float(totals['cgst']):,.2f}    SGST: INR {float(totals['sgst']):,.2f}",
        f"IGST: INR {float(totals['igst']):,.2f}    Cess: INR {float(totals['cess']):,.2f}", "", "Sales summary by category:"]
    for category in result["annexure_report"]["rows"]:
        if category["records"]:
            lines.append(f"  {category['particulars']} | {category['records']} documents | Taxable INR {float(category['taxable_amount']):,.2f} | GST INR {float(category['cgst'] + category['sgst'] + category['igst'] + category['cess']):,.2f}")
    lines.extend(["", result["annexure_report"]["note"]])
    lines.extend(["", "This is a preparation working paper. It does not confirm filing or portal reconciliation."])
    def pdf_text(value):
        safe = str(value).encode("cp1252", "replace").decode("cp1252")
        return safe.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)").replace("\r", " ").replace("\n", " ")
    # Cap each page to a readable line count and continue onto following pages
    # rather than clipping long exception lists below the print area.
    page_lines = [lines[offset:offset + 48] for offset in range(0, len(lines), 48)] or [[]]
    page_ids = [3 + index * 2 for index in range(len(page_lines))]
    font_id = 3 + len(page_lines) * 2
    objects = [b""] * font_id
    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects[1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()
    for page_index, (page_id, chunk) in enumerate(zip(page_ids, page_lines)):
        content_id = page_id + 1
        commands = ["BT", "/F1 11 Tf", "48 790 Td"]
        for index, line in enumerate(chunk):
            if index == 0 and page_index == 0:
                commands.extend(["/F1 16 Tf", f"({pdf_text(line)}) Tj", "/F1 10 Tf", "0 -25 Td"])
            else:
                commands.extend([f"({pdf_text(line)[:150]}) Tj", "0 -15 Td"])
        commands.append("ET")
        content = "\n".join(commands).encode("cp1252", "replace")
        objects[page_id - 1] = (f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>").encode()
        objects[content_id - 1] = b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream"
    objects[font_id - 1] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(pdf)); pdf.extend(f"{number} 0 obj\n".encode()); pdf.extend(body); pdf.extend(b"\nendobj\n")
    xref = len(pdf); pdf.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010} 00000 n \n".encode())
    pdf.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    audit.record(tenant.db, firm_id=tenant.firm_id, user_id=actor.id, action=AuditAction.GSTR1_DRAFT_EXPORTED,
        entity_type="GSTR1Return", entity_id=result["id"], metadata={"period": period, "format": "ANNEXURE_PDF"},
        ip_address=client_ip(request), user_agent=user_agent(request))
    tenant.commit()
    return StreamingResponse(iter([bytes(pdf)]), media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="GSTR1-{period}-annexure-summary.pdf"'})


@router.get("/gstins/{gstin_id}/returns/{period}/draft-json")
def export_draft_json(gstin_id: str, period: str, request: Request,
                      actor: models.User = Depends(require_role(Role.ASSOCIATE)),
                      tenant: TenantSession = Depends(get_tenant)):
    """Structured review export; never claims portal filing or final submission."""
    result = _get_return_data(gstin_id, period, actor, tenant)
    if result["status"] == models.GSTR1ReturnStatus.APPROVED.value:
        raise HTTPException(status_code=409, detail="This return is partner approved; use the approved preparation export.")
    audit.record(tenant.db, firm_id=tenant.firm_id, user_id=actor.id, action=AuditAction.GSTR1_DRAFT_EXPORTED,
        entity_type="GSTR1Return", entity_id=result["id"], metadata={"period": period, "document_status": "DRAFT"},
        ip_address=client_ip(request), user_agent=user_agent(request))
    tenant.commit()
    return {"document_status": "DRAFT — REVIEW REQUIRED — NOT FILED", "return": result,
            "exported_at": datetime.utcnow().isoformat() + "Z"}
